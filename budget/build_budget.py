"""בונה את טבלת ניהול התקציב השנתי (budget.xlsx) — מותאמת ל-Google Sheets.

הרצה:  python budget/build_budget.py
הקובץ נוצר ליד הסקריפט. אחרי ההרצה מומלץ לחשב את הנוסחאות (recalc) כדי שערכים
יישמרו בקובץ, ואז להעלות ל-Google Drive ולפתוח עם Google Sheets.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("budget.xlsx")
YEAR = 2026

# ---------------------------------------------------------------- styles
FONT = "Arial"
CUR = '₪#,##0;[Red]-₪#,##0;"–"'
PCT = '0%;[Red]-0%;"–"'
DATE = "dd/mm/yyyy"

DARK = "2F3E46"
GREY_TXT = "6B7280"
INPUT = "FFF7D6"      # רקע לתאים שממלאים
LINE = "D9DEE3"
BG_SOFT = "F6F8FA"

TYPES = ["הכנסות", "תוכניות חיסכון", "חשבונות ומנויים", "הוצאות משתנות", "חובות"]
# (header colour, light colour) per type — פסטל כמו בדוגמה
COLORS = {
    "הכנסות": ("9FC5E8", "E3EFFA"),
    "תוכניות חיסכון": ("93C47D", "E6F2E0"),
    "חשבונות ומנויים": ("C9A3D9", "F2E8F7"),
    "הוצאות משתנות": ("F6B26B", "FDEBD8"),
    "חובות": ("E06666", "F9DEDE"),
}
OUTFLOW = TYPES[1:]

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BOTTOM = Border(bottom=thin)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def font(size=10, bold=False, color="1F2937", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=False)
WRAP_R = Alignment(horizontal="right", vertical="top", wrap_text=True)


def q(sheet):
    return f"'{sheet}'"


def style(c, *, f=None, bg=None, fmt=None, al=None, border=None):
    if f:
        c.font = f
    if bg:
        c.fill = fill(bg)
    if fmt:
        c.number_format = fmt
    if al:
        c.alignment = al
    if border:
        c.border = border
    return c


def setup(ws, tab_color, widths):
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab_color
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def title(ws, rng, text, sub=None, sub_rng=None):
    first = rng.split(":")[0]
    ws.merge_cells(rng)
    ws[first] = text
    style(ws[first], f=font(20, True, "FFFFFF"), bg=DARK, al=CENTER)
    ws.row_dimensions[ws[first].row].height = 40
    for row in ws[rng]:
        for c in row:
            c.fill = fill(DARK)
    if sub:
        s_first = sub_rng.split(":")[0]
        ws.merge_cells(sub_rng)
        ws[s_first] = sub
        style(ws[s_first], f=font(10, False, GREY_TXT, italic=True), al=CENTER)
        ws.row_dimensions[ws[s_first].row].height = 20


def card(ws, row, c1, c2, label, formula, fmt, color):
    """כרטיס KPI: תווית בשורה row וערך גדול בשורה row+1, ממוזג על c1..c2."""
    ws.merge_cells(f"{c1}{row}:{c2}{row}")
    ws.merge_cells(f"{c1}{row+1}:{c2}{row+1}")
    style(ws[f"{c1}{row}"], f=font(10, True, DARK), al=CENTER)
    ws[f"{c1}{row}"] = label
    style(ws[f"{c1}{row+1}"], f=font(18, True, DARK), al=CENTER, fmt=fmt)
    ws[f"{c1}{row+1}"] = formula
    for r in (row, row + 1):
        for ci in range(col_idx(c1), col_idx(c2) + 1):
            cell = ws.cell(r, ci)
            cell.fill = fill(color)
    ws.row_dimensions[row].height = 20
    ws.row_dimensions[row + 1].height = 34


def pct_labels():
    return DataLabelList(showPercent=True, showVal=False, showCatName=False,
                         showSerName=False, showLegendKey=False, showLeaderLines=False)


def col_idx(letter):
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(letter)


def header_row(ws, row, c1, values, color, txt="1F2937"):
    for i, v in enumerate(values):
        c = ws.cell(row, col_idx(c1) + i, v)
        style(c, f=font(10, True, txt), bg=color, al=CENTER, border=BOX)


# ---------------------------------------------------------------- data
MONTHS = ["ינואר", "פברואר", "מרץ", "אפריל", "מאי", "יוני",
          "יולי", "אוגוסט", "ספטמבר", "אוקטובר", "נובמבר", "דצמבר"]
FREQS = [("חודשי", "=1"), ("דו-שבועי", "=26/12"), ("שבועי", "=52/12"),
         ("דו-חודשי", "=1/2"), ("שנתי", "=1/12")]

# קטגוריות ברירת מחדל: (שם, סכום לתקופה, תדירות)
CATS = {
    "הכנסות": [("משכורת 1", 14000, "חודשי"), ("משכורת 2", 9500, "חודשי"),
               ("עבודה נוספת", 1500, "חודשי"), ("קצבת ילדים", 320, "חודשי"),
               ("החזרים והכנסות אחרות", 0, "חודשי"), None, None, None],
    "תוכניות חיסכון": [("קרן חירום", 1500, "חודשי"), ("חופשה משפחתית", 600, "חודשי"),
                       ("רכב חדש", 800, "חודשי"), ("דירה", 1500, "חודשי"),
                       ("השקעות", 1000, "חודשי")],
    "חשבונות ומנויים": [("שכר דירה / משכנתא", 5200, "חודשי"), ("ארנונה", 900, "דו-חודשי"),
                         ("חשמל", 700, "דו-חודשי"), ("מים", 250, "דו-חודשי"),
                         ("גז", 60, "חודשי"), ("ועד בית", 250, "חודשי"),
                         ("אינטרנט וטלוויזיה", 180, "חודשי"), ("סלולר", 120, "חודשי"),
                         ("ביטוחי בריאות וחיים", 450, "חודשי"), ("ביטוח רכב", 4800, "שנתי"),
                         ("סטרימינג ומנויים", 90, "חודשי"), ("חדר כושר", 250, "חודשי")],
    "הוצאות משתנות": [("סופרמרקט", 650, "שבועי"), ("דלק", 600, "חודשי"),
                       ("תחבורה ציבורית וחניה", 150, "חודשי"), ("מסעדות ובילויים", 250, "שבועי"),
                       ("ביגוד והנעלה", 400, "חודשי"), ("בריאות ותרופות", 200, "חודשי"),
                       ("ילדים וחינוך", 900, "חודשי"), ("מתנות ואירועים", 300, "חודשי"),
                       ("טיפוח", 200, "חודשי"), ("בית ותחזוקה", 300, "חודשי"),
                       ("חיות מחמד", 150, "חודשי"), ("שונות", 300, "חודשי"), None, None, None],
    "חובות": [("הלוואת רכב", 1100, "חודשי"), ("הלוואה בנקאית", 800, "חודשי"),
              ("כרטיס אשראי – פריסה", 400, "חודשי"), None, None, None],
}
SAVINGS_DEFAULTS = [(30000, 18000, dt.date(2027, 6, 30)), (12000, 2500, dt.date(2027, 7, 31)),
                    (40000, 15000, dt.date(2028, 12, 31)), (150000, 62000, dt.date(2030, 12, 31)),
                    (50000, 21000, dt.date(2028, 12, 31))]
DEBT_DEFAULTS = [(60000, 38000, 0.055, 1100), (30000, 22000, 0.07, 800),
                 (6000, 4800, 0.12, 400)]

# שורות בגיליון ההגדרות
S = "הגדרות"
S_FIRST = 8
rows_by_type: dict[str, list[int]] = {}
r = S_FIRST
for t in TYPES:
    rows_by_type[t] = list(range(r, r + len(CATS[t])))
    r += len(CATS[t])
S_LAST = r - 1
CAT_RANGE = f"{q(S)}!$B${S_FIRST}:$B${S_LAST}"
TYPE_RANGE = f"{q(S)}!$A${S_FIRST}:$A${S_LAST}"
YEAR_REF = f"{q(S)}!$C$3"

T = "תנועות"
T_FIRST, T_LAST = 6, 2005
TX_DATE = f"{q(T)}!$A${T_FIRST}:$A${T_LAST}"
TX_CAT = f"{q(T)}!$B${T_FIRST}:$B${T_LAST}"
TX_TYPE = f"{q(T)}!$C${T_FIRST}:$C${T_LAST}"
TX_AMT = f"{q(T)}!$E${T_FIRST}:$E${T_LAST}"

YEAR_START = f"DATE({YEAR_REF},1,1)"
YEAR_END = f"DATE({YEAR_REF}+1,1,1)"


def in_year(crit_range, crit):
    return (f"SUMIFS({TX_AMT},{crit_range},{crit},{TX_DATE},\">=\"&{YEAR_START},"
            f"{TX_DATE},\"<\"&{YEAR_END})")


wb = Workbook()
wb.remove(wb.active)

# ================================================================ הוראות
ws_help = wb.create_sheet("הוראות")
ws_dash = wb.create_sheet("לוח שנתי")
ws_tx = wb.create_sheet(T)
month_ws = [wb.create_sheet(m) for m in MONTHS]
ws_sav = wb.create_sheet("חיסכון")
ws_debt = wb.create_sheet("חובות")
ws_set = wb.create_sheet(S)

# ================================================================ הגדרות
ws = ws_set
setup(ws, "6B7280", {"A": 18, "B": 26, "C": 16, "D": 14, "E": 16, "F": 30, "G": 3, "H": 14, "I": 14})
title(ws, "A1:I1", "⚙️ הגדרות – שנה, קטגוריות ותקציב", )
ws["A3"] = "שנת התקציב:"
style(ws["A3"], f=font(11, True), al=RIGHT)
ws["C3"] = YEAR
style(ws["C3"], f=font(12, True, "1D4ED8"), bg=INPUT, al=CENTER, border=BOX)
ws["D3"] = "← שנה כאן והכל מתעדכן"
style(ws["D3"], f=font(9, italic=True, color=GREY_TXT))
ws["A5"] = ("רשום כאן את הקטגוריות שלך. הזן סכום ותדירות (חודשי / דו-שבועי / שבועי / דו-חודשי / שנתי) "
            "והטבלה תחשב אוטומטית את התקציב החודשי. תאים צהובים = למילוי.")
ws.merge_cells("A5:F5")
style(ws["A5"], f=font(9, color=GREY_TXT), al=WRAP_R)
ws.row_dimensions[5].height = 30
header_row(ws, 7, "A", ["סוג", "קטגוריה", "סכום לתקופה", "תדירות", "תקציב חודשי", "הערות"], DARK, "FFFFFF")

dv_freq = DataValidation(type="list", formula1=f"=$H$8:$H${7+len(FREQS)}", allow_blank=True)
ws.add_data_validation(dv_freq)
for t in TYPES:
    head, light = COLORS[t]
    for i, rr in enumerate(rows_by_type[t]):
        item = CATS[t][i]
        style(ws.cell(rr, 1, t), f=font(10, True), bg=head if i == 0 else light, al=RIGHT, border=BOX)
        if i:
            ws.cell(rr, 1).font = font(9, color=GREY_TXT)
        b = style(ws.cell(rr, 2, item[0] if item else None), f=font(10), bg=INPUT, al=RIGHT, border=BOX)
        c = style(ws.cell(rr, 3, item[1] if item else None), f=font(10, color="1D4ED8"), bg=INPUT, fmt=CUR, al=CENTER, border=BOX)
        d = style(ws.cell(rr, 4, item[2] if item else "חודשי"), f=font(10), bg=INPUT, al=CENTER, border=BOX)
        dv_freq.add(d)
        ws.cell(rr, 5, f'=IF(OR(B{rr}="",C{rr}=""),0,C{rr}*IFERROR(INDEX($I$8:$I$12,MATCH(D{rr},$H$8:$H$12,0)),1))')
        style(ws.cell(rr, 5), f=font(10, True), fmt=CUR, al=CENTER, border=BOX, bg=BG_SOFT)
        style(ws.cell(rr, 6), f=font(9, color=GREY_TXT), al=RIGHT, border=BOX)
ws.cell(rows_by_type["חשבונות ומנויים"][9], 6, "תשלום שנתי – מחולק ל-12 חודשים")
ws.cell(rows_by_type["הוצאות משתנות"][0], 6, "תקציב שבועי × 52 ÷ 12")

header_row(ws, 7, "H", ["תדירות", "מקדם חודשי"], "E5E7EB")
for i, (name, factor) in enumerate(FREQS):
    style(ws.cell(8 + i, 8, name), f=font(10), al=CENTER, border=BOX)
    style(ws.cell(8 + i, 9, factor), f=font(10), fmt="0.00", al=CENTER, border=BOX)
header_row(ws, 15, "H", ["חודש", "מס'"], "E5E7EB")
for i, m in enumerate(MONTHS):
    style(ws.cell(16 + i, 8, m), f=font(10), al=CENTER, border=BOX)
    style(ws.cell(16 + i, 9, i + 1), f=font(10), al=CENTER, border=BOX)
MONTH_NAMES = f"{q(S)}!$H$16:$H$27"
ws.freeze_panes = "A8"

# ================================================================ תנועות
ws = ws_tx
setup(ws, "1F2937", {"A": 14, "B": 26, "C": 18, "D": 34, "E": 14, "F": 12, "G": 3, "H": 26, "I": 14})
title(ws, "A1:F1", "📝 יומן תנועות – כל הכנסה והוצאה במקום אחד")
ws.merge_cells("A2:F2")
ws["A2"] = ("הזן תאריך, בחר קטגוריה מהרשימה וסכום חיובי. הסוג והחודש מתמלאים לבד, "
            "וכל הלשוניות (חודשים, לוח שנתי, חיסכון וחובות) מתעדכנות אוטומטית.")
style(ws["A2"], f=font(9, color=GREY_TXT, italic=True), al=CENTER)
ws.row_dimensions[2].height = 30
header_row(ws, 5, "A", ["תאריך", "קטגוריה", "סוג (אוטומטי)", "תיאור / הערה", "סכום", "חודש"], DARK, "FFFFFF")
ws.freeze_panes = "A6"

dv_cat = DataValidation(type="list", formula1=f"={CAT_RANGE}", allow_blank=True,
                        showErrorMessage=True, error="בחר קטגוריה מהרשימה (מגיליון ההגדרות)",
                        errorTitle="קטגוריה לא קיימת")
dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                         showErrorMessage=True, error="הזן תאריך תקין, למשל 15/01/2026")
dv_amt = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        showErrorMessage=True, error="הזן סכום חיובי")
for dv in (dv_cat, dv_date, dv_amt):
    ws.add_data_validation(dv)
dv_cat.add(f"B{T_FIRST}:B{T_LAST}")
dv_date.add(f"A{T_FIRST}:A{T_LAST}")
dv_amt.add(f"E{T_FIRST}:E{T_LAST}")

for rr in range(T_FIRST, T_LAST + 1):
    band = "FFFFFF" if rr % 2 == 0 else BG_SOFT
    style(ws.cell(rr, 1), f=font(10), fmt=DATE, al=CENTER, bg=band, border=BOTTOM)
    style(ws.cell(rr, 2), f=font(10), al=RIGHT, bg=band, border=BOTTOM)
    ws.cell(rr, 3, f'=IF(B{rr}="","",IFERROR(INDEX({TYPE_RANGE},MATCH(B{rr},{CAT_RANGE},0)),"⚠ לא בהגדרות"))')
    style(ws.cell(rr, 3), f=font(9, color=GREY_TXT), al=CENTER, bg=band, border=BOTTOM)
    style(ws.cell(rr, 4), f=font(10), al=RIGHT, bg=band, border=BOTTOM)
    style(ws.cell(rr, 5), f=font(10, True), fmt=CUR, al=CENTER, bg=band, border=BOTTOM)
    ws.cell(rr, 6, f'=IF(A{rr}="","",INDEX({MONTH_NAMES},MONTH(A{rr})))')
    style(ws.cell(rr, 6), f=font(9, color=GREY_TXT), al=CENTER, bg=band, border=BOTTOM)

# צבע לפי סוג בעמודת הסוג
for t in TYPES:
    ws.conditional_formatting.add(
        f"C{T_FIRST}:C{T_LAST}",
        CellIsRule(operator="equal", formula=[f'"{t}"'], fill=fill(COLORS[t][1]), font=font(9, True)))

# סיכום מהיר ליד היומן
header_row(ws, 5, "H", ["סיכום השנה לפי סוג", "סכום"], "E5E7EB")
for i, t in enumerate(TYPES):
    style(ws.cell(6 + i, 8, t), f=font(10, True), bg=COLORS[t][1], al=RIGHT, border=BOX)
    ws.cell(6 + i, 9, f'={in_year(TX_TYPE, f"H{6+i}")}')
    style(ws.cell(6 + i, 9), f=font(10, True), fmt=CUR, al=CENTER, border=BOX)
style(ws.cell(11, 8, "תנועות שהוזנו"), f=font(10), al=RIGHT, border=BOX)
ws.cell(11, 9, f"=COUNT(E{T_FIRST}:E{T_LAST})")
style(ws.cell(11, 9), f=font(10), al=CENTER, border=BOX)

# תנועות לדוגמה – ינואר
sample = [
    (1, "משכורת 1", "משכורת ינואר", 14000), (1, "משכורת 2", "משכורת ינואר", 9500),
    (1, "שכר דירה / משכנתא", "", 5200), (2, "קרן חירום", "הוראת קבע", 1500),
    (2, "דירה", "הוראת קבע", 1500), (2, "השקעות", "קרן מחקה מדד", 1000),
    (3, "סופרמרקט", "שופרסל", 612), (4, "דלק", "", 280), (5, "סלולר", "", 120),
    (5, "אינטרנט וטלוויזיה", "", 180), (6, "מסעדות ובילויים", "ארוחת ערב", 260),
    (8, "סופרמרקט", "רמי לוי", 705), (10, "הלוואת רכב", "", 1100),
    (10, "הלוואה בנקאית", "", 800), (10, "כרטיס אשראי – פריסה", "", 400),
    (11, "ילדים וחינוך", "חוג + צהרון", 950), (12, "ביגוד והנעלה", "מבצע חורף", 520),
    (14, "חשמל", "חשבון דו-חודשי", 690), (15, "סופרמרקט", "", 580), (15, "חופשה משפחתית", "", 600),
    (15, "רכב חדש", "", 800), (17, "ועד בית", "", 250), (18, "דלק", "", 310),
    (19, "מסעדות ובילויים", "סרט + פיצה", 190), (20, "עבודה נוספת", "פרויקט פרילנס", 1800),
    (21, "סופרמרקט", "", 640), (23, "בריאות ותרופות", "בית מרקחת", 140),
    (24, "מתנות ואירועים", "חתונה", 500), (25, "סטרימינג ומנויים", "", 90),
    (25, "ביטוחי בריאות וחיים", "", 450), (27, "חדר כושר", "", 250),
    (28, "סופרמרקט", "", 690), (29, "שונות", "", 120), (30, "בית ותחזוקה", "אינסטלטור", 350),
]
for i, (day, cat, desc, amt) in enumerate(sample):
    rr = T_FIRST + i
    ws.cell(rr, 1, dt.date(YEAR, 1, day))
    ws.cell(rr, 2, cat)
    ws.cell(rr, 4, desc or None)
    ws.cell(rr, 5, amt)

# ================================================================ לשוניות חודשיות
# עמודות טבלאות הקטגוריות (5 טבלאות זו לצד זו)
TABLE_COLS = {"הכנסות": "A", "תוכניות חיסכון": "F", "חשבונות ומנויים": "K",
              "הוצאות משתנות": "P", "חובות": "U"}
M_HEAD, M_FIRST = 25, 26
M_ROWS = max(len(v) for v in CATS.values())
M_TOTAL = M_FIRST + M_ROWS
SUMROW = {t: 8 + i for i, t in enumerate(TYPES)}  # שורות טבלת הסיכום (A..D)
REMAIN_ROW = 13
planned_cells: dict[int, list[str]] = {r_: [] for t in TYPES for r_ in rows_by_type[t]}


def off(letter, n):
    return L(col_idx(letter) + n)


month_widths = {}
for t, c0 in TABLE_COLS.items():
    month_widths[c0] = 20
    for k in (1, 2, 3):
        month_widths[off(c0, k)] = 11
for sp in ("E", "J", "O", "T"):
    month_widths[sp] = 2

for mi, ws in enumerate(month_ws):
    m = mi + 1
    setup(ws, COLORS[TYPES[mi % 5]][0], month_widths)
    title(ws, "A1:X1", f'=" 🗓️ "&"{MONTHS[mi]}"&" "&{YEAR_REF}&" – תקציב חודשי"',
          "התקציב נמשך מלשונית ההגדרות (אפשר לשנות ידנית כאן לחודש הזה בלבד) • הבפועל נמשך מיומן התנועות",
          "A2:X2")
    # תאי עזר מוסתרים
    ws["Z1"] = f"=DATE({YEAR_REF},{m},1)"
    ws["Z2"] = f"=DATE({YEAR_REF},{m + 1},1)"
    ws["Z3"] = "=Z2-Z1"
    ws["Z1"].number_format = ws["Z2"].number_format = DATE
    ws.column_dimensions["Z"].hidden = True
    D_FROM, D_TO = "$Z$1", "$Z$2"

    def tx_sum(crit_range, crit):
        return (f"SUMIFS({TX_AMT},{crit_range},{crit},{TX_DATE},\">=\"&{D_FROM},"
                f"{TX_DATE},\"<\"&{D_TO})")

    # --- טבלאות קטגוריות
    for t, c0 in TABLE_COLS.items():
        head, light = COLORS[t]
        c1, c2, c3 = off(c0, 1), off(c0, 2), off(c0, 3)
        ws.merge_cells(f"{c0}{M_HEAD - 1}:{c3}{M_HEAD - 1}")
        style(ws[f"{c0}{M_HEAD - 1}"], f=font(11, True), bg=head, al=CENTER)
        ws[f"{c0}{M_HEAD - 1}"] = t
        for cc in (c1, c2, c3):
            ws[f"{cc}{M_HEAD - 1}"].fill = fill(head)
        last = "הפרש" if t == "הכנסות" else "נותר"
        header_row(ws, M_HEAD, c0, ["קטגוריה", "תקציב", "בפועל", last], light)
        srows = rows_by_type[t]
        for i in range(M_ROWS):
            rr = M_FIRST + i
            band = "FFFFFF" if i % 2 == 0 else BG_SOFT
            if i < len(srows):
                sr = srows[i]
                ws[f"{c0}{rr}"] = f'=IF({q(S)}!$B${sr}="","",{q(S)}!$B${sr})'
                ws[f"{c1}{rr}"] = f'=IF({c0}{rr}="","",{q(S)}!$E${sr})'
                ws[f"{c2}{rr}"] = f'=IF({c0}{rr}="","",{tx_sum(TX_CAT, f"{c0}{rr}")})'
                if t == "הכנסות":
                    ws[f"{c3}{rr}"] = f'=IF({c0}{rr}="","",N({c2}{rr})-N({c1}{rr}))'
                else:
                    ws[f"{c3}{rr}"] = f'=IF({c0}{rr}="","",N({c1}{rr})-N({c2}{rr}))'
                planned_cells[sr].append(f"{q(MONTHS[mi])}!{c1}{rr}")
            style(ws[f"{c0}{rr}"], f=font(9), al=RIGHT, bg=band, border=BOX)
            style(ws[f"{c1}{rr}"], f=font(9, color="1D4ED8"), fmt=CUR, al=CENTER, bg=band, border=BOX)
            style(ws[f"{c2}{rr}"], f=font(9, True), fmt=CUR, al=CENTER, bg=band, border=BOX)
            style(ws[f"{c3}{rr}"], f=font(9), fmt=CUR, al=CENTER, bg=band, border=BOX)
        ws[f"{c0}{M_TOTAL}"] = 'סה"כ'
        for cc in (c1, c2, c3):
            ws[f"{cc}{M_TOTAL}"] = f"=SUM({cc}{M_FIRST}:{cc}{M_TOTAL - 1})"
        for cc in (c0, c1, c2, c3):
            style(ws[f"{cc}{M_TOTAL}"], f=font(10, True), bg=head, fmt=CUR, al=CENTER, border=BOX)
        # חריגה מהתקציב בהוצאות – אדום
        if t != "הכנסות":
            ws.conditional_formatting.add(
                f"{c2}{M_FIRST}:{c2}{M_TOTAL}",
                FormulaRule(formula=[f'AND(ISNUMBER({c2}{M_FIRST}),{c2}{M_FIRST}>{c1}{M_FIRST})'],
                            font=Font(name=FONT, color="C00000", bold=True), fill=fill("FDE2E2")))

    # --- טבלת סיכום A7:D13
    header_row(ws, 7, "A", ["סיכום החודש", "תקציב", "בפועל", "הפרש"], "E5E7EB")
    for t in TYPES:
        rr = SUMROW[t]
        c0 = TABLE_COLS[t]
        style(ws.cell(rr, 1, t), f=font(10, True), bg=COLORS[t][1], al=RIGHT, border=BOX)
        ws.cell(rr, 2, f"={off(c0, 1)}{M_TOTAL}")
        ws.cell(rr, 3, f"={off(c0, 2)}{M_TOTAL}")
        ws.cell(rr, 4, f"=C{rr}-B{rr}" if t == "הכנסות" else f"=B{rr}-C{rr}")
        for cc in (2, 3, 4):
            style(ws.cell(rr, cc), f=font(10, cc == 3), fmt=CUR, al=CENTER, border=BOX)
    style(ws.cell(REMAIN_ROW, 1, "נשאר בעו\"ש"), f=font(10, True, "FFFFFF"), bg=DARK, al=RIGHT, border=BOX)
    ws.cell(REMAIN_ROW, 2, "=B8-SUM(B9:B12)")
    ws.cell(REMAIN_ROW, 3, "=C8-SUM(C9:C12)")
    ws.cell(REMAIN_ROW, 4, "=C13-B13")
    for cc in (2, 3, 4):
        style(ws.cell(REMAIN_ROW, cc), f=font(10, True, "FFFFFF"), bg=DARK, fmt=CUR, al=CENTER, border=BOX)

    # --- מעקב שבועי להוצאות משתנות F7:I13
    header_row(ws, 7, "F", ["שבוע (משתנות)", "תקציב", "בפועל", "נותר"], COLORS["הוצאות משתנות"][1])
    weeks = [("ימים 1–7", 0, 7), ("ימים 8–14", 7, 14), ("ימים 15–21", 14, 21),
             ("ימים 22–28", 21, 28), ("ימים 29–סוף", 28, None)]
    for i, (lbl, a, b) in enumerate(weeks):
        rr = 8 + i
        days = "($Z$3-28)" if b is None else "7"
        end = "$Z$2" if b is None else f"$Z$1+{b}"
        style(ws.cell(rr, 6, lbl), f=font(10), al=RIGHT, border=BOX)
        ws.cell(rr, 7, f"=$B$11*{days}/$Z$3")
        ws.cell(rr, 8, f'=SUMIFS({TX_AMT},{TX_TYPE},"הוצאות משתנות",{TX_DATE},">="&($Z$1+{a}),{TX_DATE},"<"&{end})')
        ws.cell(rr, 9, f"=G{rr}-H{rr}")
        for cc in (7, 8, 9):
            style(ws.cell(rr, cc), f=font(10, cc == 8), fmt=CUR, al=CENTER, border=BOX)
    style(ws.cell(13, 6, 'סה"כ'), f=font(10, True), bg=COLORS["הוצאות משתנות"][0], al=RIGHT, border=BOX)
    for cc, col in ((7, "G"), (8, "H"), (9, "I")):
        ws.cell(13, cc, f"=SUM({col}8:{col}12)")
        style(ws.cell(13, cc), f=font(10, True), bg=COLORS["הוצאות משתנות"][0], fmt=CUR, al=CENTER, border=BOX)
    ws.merge_cells("A15:I16")
    ws["A15"] = ("💡 תקציב שבועי להוצאות משתנות = התקציב החודשי מחולק לפי מספר הימים בכל שבוע. "
                 "אדום = חריגה מהתקציב.")
    style(ws["A15"], f=font(9, color=GREY_TXT, italic=True), al=WRAP_R)

    # --- כרטיסי KPI
    card(ws, 4, "A", "D", "💵 הכנסות החודש", "=C8", CUR, COLORS["הכנסות"][1])
    card(ws, 4, "F", "I", "💸 סה\"כ יציאות", "=SUM(C9:C12)", CUR, COLORS["חובות"][1])
    card(ws, 4, "K", "N", "🏦 נשאר בעו\"ש", "=C13", CUR, "E5E7EB")
    card(ws, 4, "P", "S", "🐷 שיעור חיסכון", '=IF(C8=0,"–",C9/C8)', PCT, COLORS["תוכניות חיסכון"][1])
    card(ws, 4, "U", "X", "🛒 ניצול תקציב משתנות", '=IF(B11=0,"–",C11/B11)', PCT, COLORS["הוצאות משתנות"][1])

    # --- גרפים
    pie = PieChart()
    pie.title = "לאן הלך הכסף החודש"
    pie.add_data(Reference(ws, min_col=3, min_row=9, max_row=12), titles_from_data=False)
    pie.set_categories(Reference(ws, min_col=1, min_row=9, max_row=12))
    for i, t in enumerate(OUTFLOW):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = COLORS[t][0]
        pie.series[0].dPt.append(pt)
    pie.dataLabels = pct_labels()
    pie.width, pie.height = 10, 7.6
    ws.add_chart(pie, "K7")

    bar = BarChart()
    bar.type = "col"
    bar.title = "תקציב מול בפועל"
    bar.add_data(Reference(ws, min_col=2, max_col=3, min_row=7, max_row=12), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=8, max_row=12))
    bar.series[0].graphicalProperties.solidFill = "CBD5E1"
    bar.series[1].graphicalProperties.solidFill = "4A90D9"
    bar.y_axis.numFmt = "#,##0"
    bar.y_axis.majorGridlines = None
    bar.width, bar.height = 19, 7.6
    bar.legend.position = "b"
    ws.add_chart(bar, "P7")
    ws.freeze_panes = "A3"

# ================================================================ לוח שנתי
ws = ws_dash
setup(ws, DARK, {"A": 22, "B": 15, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "H": 15, "I": 3})
title(ws, "A1:H1", f'="📊 לוח שנתי – סיכום "&{YEAR_REF}',
      "סקירה מלאה של ההכנסות וההוצאות לאורך כל השנה • מתעדכן אוטומטית", "A2:H2")
D_HEAD = 7
header_row(ws, D_HEAD, "A", ["חודש"] + TYPES + ['סה"כ יציאות', "נשאר בעו\"ש"], DARK, "FFFFFF")
for i, t in enumerate(TYPES):
    ws.cell(D_HEAD, 2 + i).fill = fill(COLORS[t][0])
    ws.cell(D_HEAD, 2 + i).font = font(10, True)
for mi, mname in enumerate(MONTHS):
    rr = D_HEAD + 1 + mi
    band = "FFFFFF" if mi % 2 == 0 else BG_SOFT
    style(ws.cell(rr, 1, mname), f=font(10, True), al=RIGHT, bg=band, border=BOX)
    for i, t in enumerate(TYPES):
        ws.cell(rr, 2 + i, f"={q(mname)}!$C${SUMROW[t]}")
    ws.cell(rr, 7, f"=SUM(C{rr}:F{rr})")
    ws.cell(rr, 8, f"=B{rr}-G{rr}")
    for cc in range(2, 9):
        style(ws.cell(rr, cc), f=font(10, cc >= 7), fmt=CUR, al=CENTER, bg=band, border=BOX)
D_TOT = D_HEAD + 13
D_PLAN = D_TOT + 1
D_AVG = D_TOT + 2
style(ws.cell(D_TOT, 1, 'סה"כ בפועל'), f=font(10, True, "FFFFFF"), bg=DARK, al=RIGHT, border=BOX)
style(ws.cell(D_PLAN, 1, "תקציב שנתי"), f=font(10, True), bg="E5E7EB", al=RIGHT, border=BOX)
style(ws.cell(D_AVG, 1, "ממוצע חודשי"), f=font(10, True), bg="E5E7EB", al=RIGHT, border=BOX)
for cc in range(2, 9):
    col = L(cc)
    ws.cell(D_TOT, cc, f"=SUM({col}{D_HEAD + 1}:{col}{D_HEAD + 12})")
    style(ws.cell(D_TOT, cc), f=font(10, True, "FFFFFF"), bg=DARK, fmt=CUR, al=CENTER, border=BOX)
    if cc <= 6:
        t = TYPES[cc - 2]
        ws.cell(D_PLAN, cc, "=" + "+".join(f"{q(mn)}!$B${SUMROW[t]}" for mn in MONTHS))
    elif cc == 7:
        ws.cell(D_PLAN, cc, f"=SUM(C{D_PLAN}:F{D_PLAN})")
    else:
        ws.cell(D_PLAN, cc, f"=B{D_PLAN}-G{D_PLAN}")
    ws.cell(D_AVG, cc, f"={col}{D_TOT}/12")
    for r_ in (D_PLAN, D_AVG):
        style(ws.cell(r_, cc), f=font(10), bg="F3F4F6", fmt=CUR, al=CENTER, border=BOX)

card(ws, 4, "A", "B", "💵 הכנסות השנה", f"=B{D_TOT}", CUR, COLORS["הכנסות"][1])
card(ws, 4, "C", "D", "💸 יציאות השנה", f"=G{D_TOT}", CUR, COLORS["חובות"][1])
card(ws, 4, "E", "F", "🏦 נשאר בעו\"ש", f"=H{D_TOT}", CUR, "E5E7EB")
card(ws, 4, "G", "H", "🐷 שיעור חיסכון", f'=IF(B{D_TOT}=0,"–",C{D_TOT}/B{D_TOT})', PCT, COLORS["תוכניות חיסכון"][1])

# פירוט לפי קטגוריה
C_HEAD = D_AVG + 4
ws.merge_cells(f"A{C_HEAD - 1}:F{C_HEAD - 1}")
ws[f"A{C_HEAD - 1}"] = "📋 פירוט שנתי לפי קטגוריה"
style(ws[f"A{C_HEAD - 1}"], f=font(12, True, DARK), al=RIGHT)
header_row(ws, C_HEAD, "A", ["קטגוריה", "סוג", "תקציב שנתי", "בפועל", "הפרש", "% ניצול"], "E5E7EB")
rr = C_HEAD
for t in TYPES:
    for sr in rows_by_type[t]:
        rr += 1
        ws.cell(rr, 1, f'=IF({q(S)}!$B${sr}="","",{q(S)}!$B${sr})')
        ws.cell(rr, 2, f'=IF(A{rr}="","","{t}")')
        ws.cell(rr, 3, f'=IF(A{rr}="","",' + "+".join(f"N({c})" for c in planned_cells[sr]) + ")")
        ws.cell(rr, 4, f'=IF(A{rr}="","",{in_year(TX_CAT, f"A{rr}")})')
        ws.cell(rr, 5, f'=IF(A{rr}="","",D{rr}-C{rr})' if t == "הכנסות" else f'=IF(A{rr}="","",C{rr}-D{rr})')
        ws.cell(rr, 6, f'=IF(OR(A{rr}="",N(C{rr})=0),"",D{rr}/C{rr})')
        style(ws.cell(rr, 1), f=font(10), al=RIGHT, border=BOX)
        style(ws.cell(rr, 2), f=font(9, color=GREY_TXT), bg=COLORS[t][1], al=CENTER, border=BOX)
        for cc in (3, 4, 5):
            style(ws.cell(rr, cc), f=font(10, cc == 4), fmt=CUR, al=CENTER, border=BOX)
        style(ws.cell(rr, 6), f=font(10), fmt=PCT, al=CENTER, border=BOX)
C_LAST = rr
ws.conditional_formatting.add(
    f"F{C_HEAD + 1}:F{C_LAST}",
    FormulaRule(formula=[f'AND($B{C_HEAD + 1}<>"הכנסות",ISNUMBER($F{C_HEAD + 1}),$F{C_HEAD + 1}>1)'],
                font=Font(name=FONT, color="C00000", bold=True), fill=fill("FDE2E2")))

# גרפים
line = LineChart()
line.title = "הכנסות מול יציאות לאורך השנה"
line.add_data(Reference(ws, min_col=2, min_row=D_HEAD, max_row=D_HEAD + 12), titles_from_data=True)
line.add_data(Reference(ws, min_col=7, min_row=D_HEAD, max_row=D_HEAD + 12), titles_from_data=True)
line.set_categories(Reference(ws, min_col=1, min_row=D_HEAD + 1, max_row=D_HEAD + 12))
line.series[0].graphicalProperties.line.solidFill = "4A90D9"
line.series[1].graphicalProperties.line.solidFill = "E06666"
for s in line.series:
    s.graphicalProperties.line.width = 28000
    s.smooth = True
line.x_axis.scaling.orientation = "maxMin"
line.y_axis.numFmt = "#,##0"
line.y_axis.majorGridlines = None
line.legend.position = "b"
line.width, line.height = 20, 8.5
ws.add_chart(line, "J3")

stack = BarChart()
stack.type = "col"
stack.grouping = "stacked"
stack.overlap = 100
stack.title = "התפלגות היציאות בכל חודש"
stack.add_data(Reference(ws, min_col=3, max_col=6, min_row=D_HEAD, max_row=D_HEAD + 12), titles_from_data=True)
stack.set_categories(Reference(ws, min_col=1, min_row=D_HEAD + 1, max_row=D_HEAD + 12))
for s, t in zip(stack.series, OUTFLOW):
    s.graphicalProperties.solidFill = COLORS[t][0]
stack.x_axis.scaling.orientation = "maxMin"
stack.y_axis.numFmt = "#,##0"
stack.y_axis.majorGridlines = None
stack.legend.position = "b"
stack.width, stack.height = 20, 8.5
ws.add_chart(stack, "J21")

pie = PieChart()
pie.title = "חלוקת היציאות השנתית"
pie.add_data(Reference(ws, min_row=D_TOT, min_col=3, max_col=6), from_rows=True, titles_from_data=False)
pie.set_categories(Reference(ws, min_row=D_HEAD, min_col=3, max_col=6))
for i, t in enumerate(OUTFLOW):
    pt = DataPoint(idx=i)
    pt.graphicalProperties.solidFill = COLORS[t][0]
    pie.series[0].dPt.append(pt)
pie.dataLabels = pct_labels()
pie.width, pie.height = 20, 8.5
ws.add_chart(pie, "J39")
ws.freeze_panes = "A3"

# ================================================================ חיסכון
ws = ws_sav
setup(ws, COLORS["תוכניות חיסכון"][0],
      {"A": 20, "B": 13, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13, "H": 11, "I": 16, "J": 11, "K": 13, "L": 3})
title(ws, "A1:K1", "🐷 מעקב תוכניות חיסכון",
      "עקבו אחרי 5 תוכניות חיסכון – כמה נחסך, כמה נשאר ומה צריך להפקיד כל חודש כדי לעמוד ביעד", "A2:K2")
SV_HEAD = 7
header_row(ws, SV_HEAD, "A", ["תוכנית", "יעד", "יתרת פתיחה", "תאריך יעד", "הופקד השנה",
                               'סה"כ נחסך', "נותר לחסוך", "% התקדמות", "מד התקדמות",
                               "חודשים ליעד", "נדרש לחודש"], COLORS["תוכניות חיסכון"][0])
for i, sr in enumerate(rows_by_type["תוכניות חיסכון"]):
    rr = SV_HEAD + 1 + i
    goal, opening, tdate = SAVINGS_DEFAULTS[i]
    ws.cell(rr, 1, f'=IF({q(S)}!$B${sr}="","",{q(S)}!$B${sr})')
    ws.cell(rr, 2, goal)
    ws.cell(rr, 3, opening)
    ws.cell(rr, 4, tdate)
    ws.cell(rr, 5, f'=IF(A{rr}="",0,{in_year(TX_CAT, f"A{rr}")})')
    ws.cell(rr, 6, f'=IF(A{rr}="",0,N(C{rr})+E{rr})')
    ws.cell(rr, 7, f'=IF(A{rr}="",0,MAX(0,N(B{rr})-F{rr}))')
    ws.cell(rr, 8, f'=IF(OR(A{rr}="",N(B{rr})=0),"",MIN(1,F{rr}/B{rr}))')
    ws.cell(rr, 9, f'=IF(H{rr}="","",REPT("█",ROUND(H{rr}*10,0))&REPT("░",10-ROUND(H{rr}*10,0)))')
    ws.cell(rr, 10, f'=IF(OR(A{rr}="",D{rr}=""),"",MAX(0,(YEAR(D{rr})-YEAR(TODAY()))*12+MONTH(D{rr})-MONTH(TODAY())))')
    ws.cell(rr, 11, f'=IF(J{rr}="","",IF(G{rr}<=0,0,G{rr}/MAX(1,J{rr})))')
    style(ws.cell(rr, 1), f=font(10, True), bg=COLORS["תוכניות חיסכון"][1], al=RIGHT, border=BOX)
    for cc in (2, 3):
        style(ws.cell(rr, cc), f=font(10, color="1D4ED8"), bg=INPUT, fmt=CUR, al=CENTER, border=BOX)
    style(ws.cell(rr, 4), f=font(10, color="1D4ED8"), bg=INPUT, fmt=DATE, al=CENTER, border=BOX)
    for cc in (5, 6, 7, 11):
        style(ws.cell(rr, cc), f=font(10, cc == 6), fmt=CUR, al=CENTER, border=BOX)
    style(ws.cell(rr, 8), f=font(10, True), fmt="0%", al=CENTER, border=BOX)
    style(ws.cell(rr, 9), f=Font(name=FONT, size=10, color="6AA84F"), al=CENTER, border=BOX)
    style(ws.cell(rr, 10), f=font(10), fmt="0", al=CENTER, border=BOX)
SV_TOT = SV_HEAD + 6
style(ws.cell(SV_TOT, 1, 'סה"כ'), f=font(10, True, "FFFFFF"), bg=DARK, al=RIGHT, border=BOX)
for cc in range(2, 12):
    col = L(cc)
    if cc in (2, 3, 5, 6, 7, 11):
        ws.cell(SV_TOT, cc, f"=SUM({col}{SV_HEAD + 1}:{col}{SV_HEAD + 5})")
    elif cc == 8:
        ws.cell(SV_TOT, cc, f'=IF(B{SV_TOT}=0,"",MIN(1,F{SV_TOT}/B{SV_TOT}))')
    elif cc == 9:
        ws.cell(SV_TOT, cc, f'=IF(H{SV_TOT}="","",REPT("█",ROUND(H{SV_TOT}*10,0))&REPT("░",10-ROUND(H{SV_TOT}*10,0)))')
    style(ws.cell(SV_TOT, cc), f=font(10, True, "FFFFFF"), bg=DARK,
          fmt="0%" if cc == 8 else CUR, al=CENTER, border=BOX)

card(ws, 4, "A", "B", "🎯 סך כל היעדים", f"=B{SV_TOT}", CUR, COLORS["תוכניות חיסכון"][1])
card(ws, 4, "C", "E", "💰 נחסך עד כה", f"=F{SV_TOT}", CUR, COLORS["הכנסות"][1])
card(ws, 4, "F", "H", "⏳ נותר לחסוך", f"=G{SV_TOT}", CUR, COLORS["הוצאות משתנות"][1])
card(ws, 4, "I", "K", "📈 התקדמות כוללת", f"=H{SV_TOT}", "0%", "E5E7EB")

# הפקדות לפי חודש
SM_HEAD = SV_TOT + 3
ws.merge_cells(f"A{SM_HEAD - 1}:G{SM_HEAD - 1}")
ws[f"A{SM_HEAD - 1}"] = "📅 הפקדות לפי חודש"
style(ws[f"A{SM_HEAD - 1}"], f=font(12, True, DARK), al=RIGHT)
style(ws.cell(SM_HEAD, 1, "חודש"), f=font(10, True), bg=COLORS["תוכניות חיסכון"][0], al=CENTER, border=BOX)
for i in range(5):
    ws.cell(SM_HEAD, 2 + i, f"=A{SV_HEAD + 1 + i}")
    style(ws.cell(SM_HEAD, 2 + i), f=font(10, True), bg=COLORS["תוכניות חיסכון"][0], al=CENTER, border=BOX)
style(ws.cell(SM_HEAD, 7, 'סה"כ'), f=font(10, True), bg=COLORS["תוכניות חיסכון"][0], al=CENTER, border=BOX)
for mi, mname in enumerate(MONTHS):
    rr = SM_HEAD + 1 + mi
    style(ws.cell(rr, 1, mname), f=font(10, True), al=RIGHT, border=BOX)
    for i in range(5):
        col = L(2 + i)
        ws.cell(rr, 2 + i, f'=IF({col}${SM_HEAD}="",0,SUMIFS({TX_AMT},{TX_CAT},{col}${SM_HEAD},'
                           f'{TX_DATE},">="&DATE({YEAR_REF},{mi + 1},1),{TX_DATE},"<"&DATE({YEAR_REF},{mi + 2},1)))')
        style(ws.cell(rr, 2 + i), f=font(10), fmt=CUR, al=CENTER, border=BOX)
    ws.cell(rr, 7, f"=SUM(B{rr}:F{rr})")
    style(ws.cell(rr, 7), f=font(10, True), fmt=CUR, al=CENTER, border=BOX)

bar = BarChart()
bar.type = "bar"
bar.title = "יעד מול נחסך"
bar.add_data(Reference(ws, min_col=2, min_row=SV_HEAD, max_row=SV_HEAD + 5), titles_from_data=True)
bar.add_data(Reference(ws, min_col=6, min_row=SV_HEAD, max_row=SV_HEAD + 5), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=1, min_row=SV_HEAD + 1, max_row=SV_HEAD + 5))
bar.series[0].graphicalProperties.solidFill = "E5E7EB"
bar.series[1].graphicalProperties.solidFill = "6AA84F"
bar.x_axis.numFmt = "#,##0"
bar.y_axis.numFmt = "#,##0"
bar.y_axis.majorGridlines = None
bar.legend.position = "b"
bar.width, bar.height = 16, 8
ws.add_chart(bar, "M3")

ws["N40"] = "נחסך"
ws["O40"] = f"=F{SV_TOT}"
ws["N41"] = "נותר"
ws["O41"] = f"=G{SV_TOT}"
for c in ("N40", "O40", "N41", "O41"):
    ws[c].font = font(8, color="FFFFFF")  # נתוני עזר לגרף
donut = DoughnutChart()
donut.title = "התקדמות כוללת לעבר כל היעדים"
donut.add_data(Reference(ws, min_col=15, min_row=40, max_row=41), titles_from_data=False)
donut.set_categories(Reference(ws, min_col=14, min_row=40, max_row=41))
for i, c in enumerate(("6AA84F", "E5E7EB")):
    pt = DataPoint(idx=i)
    pt.graphicalProperties.solidFill = c
    donut.series[0].dPt.append(pt)
donut.dataLabels = pct_labels()
donut.holeSize = 55
donut.width, donut.height = 16, 8
ws.add_chart(donut, "M21")

# ================================================================ חובות
ws = ws_debt
setup(ws, COLORS["חובות"][0],
      {"A": 22, "B": 13, "C": 14, "D": 10, "E": 13, "F": 13, "G": 13, "H": 10, "I": 16, "J": 12, "K": 14, "L": 3})
title(ws, "A1:K1", "💳 מעקב וסילוק חובות",
      "ראו איך כל חוב יורד אחד אחרי השני ומתי תהיו חופשיים מחובות", "A2:K2")
DB_HEAD = 7
header_row(ws, DB_HEAD, "A", ["חוב", "סכום מקורי", "יתרה בתחילת השנה", "ריבית שנתית", "החזר חודשי",
                               "שולם השנה", "יתרה נוכחית", "% נפרע", "מד התקדמות",
                               "חודשים לסיום", "סיום משוער"], COLORS["חובות"][0])
debt_rows = rows_by_type["חובות"]
for i, sr in enumerate(debt_rows):
    rr = DB_HEAD + 1 + i
    ws.cell(rr, 1, f'=IF({q(S)}!$B${sr}="","",{q(S)}!$B${sr})')
    if i < len(DEBT_DEFAULTS):
        orig, start, rate, pay = DEBT_DEFAULTS[i]
        ws.cell(rr, 2, orig)
        ws.cell(rr, 3, start)
        ws.cell(rr, 4, rate)
    ws.cell(rr, 5, f'=IF(A{rr}="","",{q(S)}!$E${sr})')
    ws.cell(rr, 6, f'=IF(A{rr}="",0,{in_year(TX_CAT, f"A{rr}")})')
    ws.cell(rr, 7, f'=IF(A{rr}="",0,MAX(0,N(C{rr})-F{rr}))')
    ws.cell(rr, 8, f'=IF(OR(A{rr}="",N(B{rr})=0),"",MAX(0,MIN(1,(B{rr}-G{rr})/B{rr})))')
    ws.cell(rr, 9, f'=IF(H{rr}="","",REPT("█",ROUND(H{rr}*10,0))&REPT("░",10-ROUND(H{rr}*10,0)))')
    ws.cell(rr, 10, f'=IF(OR(A{rr}="",N(E{rr})=0),"",IF(G{rr}<=0,0,'
                    f'IFERROR(ROUNDUP(NPER(N(D{rr})/12,-E{rr},G{rr}),0),"לא מכסה ריבית")))')
    ws.cell(rr, 11, f'=IF(ISNUMBER(J{rr}),EDATE(TODAY(),J{rr}),"")')
    style(ws.cell(rr, 1), f=font(10, True), bg=COLORS["חובות"][1], al=RIGHT, border=BOX)
    for cc in (2, 3):
        style(ws.cell(rr, cc), f=font(10, color="1D4ED8"), bg=INPUT, fmt=CUR, al=CENTER, border=BOX)
    style(ws.cell(rr, 4), f=font(10, color="1D4ED8"), bg=INPUT, fmt="0.0%", al=CENTER, border=BOX)
    for cc in (5, 6, 7):
        style(ws.cell(rr, cc), f=font(10, cc == 7), fmt=CUR, al=CENTER, border=BOX)
    style(ws.cell(rr, 8), f=font(10, True), fmt="0%", al=CENTER, border=BOX)
    style(ws.cell(rr, 9), f=Font(name=FONT, size=10, color="CC0000"), al=CENTER, border=BOX)
    style(ws.cell(rr, 10), f=font(10), fmt="0", al=CENTER, border=BOX)
    style(ws.cell(rr, 11), f=font(10), fmt="mm/yyyy", al=CENTER, border=BOX)
DB_TOT = DB_HEAD + len(debt_rows) + 1
style(ws.cell(DB_TOT, 1, 'סה"כ'), f=font(10, True, "FFFFFF"), bg=DARK, al=RIGHT, border=BOX)
for cc in range(2, 12):
    col = L(cc)
    if cc in (2, 3, 5, 6, 7):
        ws.cell(DB_TOT, cc, f"=SUM({col}{DB_HEAD + 1}:{col}{DB_TOT - 1})")
    elif cc == 8:
        ws.cell(DB_TOT, cc, f'=IF(B{DB_TOT}=0,"",MAX(0,MIN(1,(B{DB_TOT}-G{DB_TOT})/B{DB_TOT})))')
    elif cc == 9:
        ws.cell(DB_TOT, cc, f'=IF(H{DB_TOT}="","",REPT("█",ROUND(H{DB_TOT}*10,0))&REPT("░",10-ROUND(H{DB_TOT}*10,0)))')
    elif cc == 10:
        ws.cell(DB_TOT, cc, f'=IF(COUNT(J{DB_HEAD + 1}:J{DB_TOT - 1})=0,"",MAX(J{DB_HEAD + 1}:J{DB_TOT - 1}))')
    elif cc == 11:
        ws.cell(DB_TOT, cc, f'=IF(ISNUMBER(J{DB_TOT}),EDATE(TODAY(),J{DB_TOT}),"")')
    style(ws.cell(DB_TOT, cc), f=font(10, True, "FFFFFF"), bg=DARK,
          fmt={8: "0%", 10: "0", 11: "mm/yyyy"}.get(cc, CUR), al=CENTER, border=BOX)

card(ws, 4, "A", "B", "💳 סך החובות (מקורי)", f"=B{DB_TOT}", CUR, COLORS["חובות"][1])
card(ws, 4, "C", "E", "✅ שולם השנה", f"=F{DB_TOT}", CUR, COLORS["תוכניות חיסכון"][1])
card(ws, 4, "F", "H", "⏳ יתרה לסילוק", f"=G{DB_TOT}", CUR, COLORS["הוצאות משתנות"][1])
card(ws, 4, "I", "K", "🏁 חופשי מחובות ב-", f"=K{DB_TOT}", "mm/yyyy", "E5E7EB")

DM_HEAD = DB_TOT + 3
ws.merge_cells(f"A{DM_HEAD - 1}:H{DM_HEAD - 1}")
ws[f"A{DM_HEAD - 1}"] = "📅 החזרים לפי חודש"
style(ws[f"A{DM_HEAD - 1}"], f=font(12, True, DARK), al=RIGHT)
style(ws.cell(DM_HEAD, 1, "חודש"), f=font(10, True), bg=COLORS["חובות"][0], al=CENTER, border=BOX)
nd = len(debt_rows)
for i in range(nd):
    ws.cell(DM_HEAD, 2 + i, f"=A{DB_HEAD + 1 + i}")
    style(ws.cell(DM_HEAD, 2 + i), f=font(10, True), bg=COLORS["חובות"][0], al=CENTER, border=BOX)
style(ws.cell(DM_HEAD, 2 + nd, 'סה"כ'), f=font(10, True), bg=COLORS["חובות"][0], al=CENTER, border=BOX)
for mi, mname in enumerate(MONTHS):
    rr = DM_HEAD + 1 + mi
    style(ws.cell(rr, 1, mname), f=font(10, True), al=RIGHT, border=BOX)
    for i in range(nd):
        col = L(2 + i)
        ws.cell(rr, 2 + i, f'=IF({col}${DM_HEAD}="",0,SUMIFS({TX_AMT},{TX_CAT},{col}${DM_HEAD},'
                           f'{TX_DATE},">="&DATE({YEAR_REF},{mi + 1},1),{TX_DATE},"<"&DATE({YEAR_REF},{mi + 2},1)))')
        style(ws.cell(rr, 2 + i), f=font(10), fmt=CUR, al=CENTER, border=BOX)
    ws.cell(rr, 2 + nd, f"=SUM(B{rr}:{L(1 + nd)}{rr})")
    style(ws.cell(rr, 2 + nd), f=font(10, True), fmt=CUR, al=CENTER, border=BOX)
NOTE = DM_HEAD + 14
ws.merge_cells(f"A{NOTE}:K{NOTE + 2}")
ws[f"A{NOTE}"] = ("💡 טיפ לסילוק מהיר: שיטת 'כדור השלג' – משלמים מינימום בכל החובות ומפנים כל שקל נוסף לחוב הקטן ביותר. "
                  "שיטת 'המפולת' – מפנים את התוספת לחוב עם הריבית הגבוהה ביותר (חוסך הכי הרבה כסף). "
                  "שימו לב: 'יתרה נוכחית' = יתרת תחילת השנה פחות ההחזרים, ללא חישוב ריבית. "
                  "'חודשים לסיום' כן מחושב לפי הריבית וההחזר החודשי.")
style(ws[f"A{NOTE}"], f=font(9, color=GREY_TXT, italic=True), al=WRAP_R)

dbar = BarChart()
dbar.type = "bar"
dbar.grouping = "stacked"
dbar.overlap = 100
dbar.title = "שולם מול נותר לכל חוב"
dbar.add_data(Reference(ws, min_col=6, max_col=7, min_row=DB_HEAD, max_row=DB_TOT - 1), titles_from_data=True)
dbar.set_categories(Reference(ws, min_col=1, min_row=DB_HEAD + 1, max_row=DB_TOT - 1))
dbar.series[0].graphicalProperties.solidFill = "6AA84F"
dbar.series[1].graphicalProperties.solidFill = "E06666"
dbar.y_axis.numFmt = "#,##0"
dbar.y_axis.majorGridlines = None
dbar.legend.position = "b"
dbar.width, dbar.height = 16, 8
ws.add_chart(dbar, "M3")

dline = BarChart()
dline.type = "col"
dline.title = "החזרי חובות לפי חודש"
dline.add_data(Reference(ws, min_col=2 + nd, min_row=DM_HEAD, max_row=DM_HEAD + 12), titles_from_data=True)
dline.set_categories(Reference(ws, min_col=1, min_row=DM_HEAD + 1, max_row=DM_HEAD + 12))
dline.series[0].graphicalProperties.solidFill = "E06666"
dline.x_axis.scaling.orientation = "maxMin"
line.y_axis.numFmt = "#,##0"
dline.y_axis.majorGridlines = None
dline.legend = None
dline.width, dline.height = 16, 8
ws.add_chart(dline, "M21")

# ================================================================ הוראות
ws = ws_help
setup(ws, "F1C232", {"A": 4, "B": 30, "C": 80})
title(ws, "A1:C1", "📖 תבנית ניהול תקציב שנתי – איך מתחילים?",
      "טבלת תקציב בעברית מלאה • עובדת ישירות ב-Google Sheets • בלי התקנות ובלי ניסיון קודם", "A2:C2")
steps = [
    ("⚙️ שלב 1 – הגדרות", "בלשונית 'הגדרות' בחרו את שנת התקציב, ורשמו את הקטגוריות שלכם בכל אחד מ-5 הסוגים: "
     "הכנסות, תוכניות חיסכון, חשבונות ומנויים, הוצאות משתנות וחובות. לכל קטגוריה הזינו סכום ותדירות – "
     "חודשי, דו-שבועי, שבועי, דו-חודשי או שנתי – והתקציב החודשי יחושב אוטומטית."),
    ("📝 שלב 2 – יומן תנועות", "בכל פעם שנכנס או יוצא כסף, הוסיפו שורה בלשונית 'תנועות': תאריך, קטגוריה מהרשימה, "
     "תיאור (לא חובה) וסכום חיובי. הסוג והחודש מתמלאים לבד. זה המקום היחיד שבו מזינים נתונים שוטפים."),
    ("🗓️ שלב 3 – לשוניות החודשים", "12 לשוניות (ינואר–דצמבר) מציגות לכל חודש: כרטיסי סיכום, תקציב מול בפועל לכל קטגוריה, "
     "מעקב שבועי להוצאות המשתנות וגרפים. רוצים תקציב שונה לחודש מסוים? פשוט הקלידו סכום בעמודת 'תקציב' באותו חודש."),
    ("📊 שלב 4 – לוח שנתי", "תמונה מלאה של כל השנה: הכנסות מול יציאות לפי חודש, תקציב שנתי מול בפועל, "
     "פירוט לכל קטגוריה וגרפים שמראים בדיוק לאן הכסף הולך."),
    ("🐷 שלב 5 – חיסכון", "הזינו לכל תוכנית יעד, יתרת פתיחה ותאריך יעד. ההפקדות נמשכות מהתנועות, "
     "ותראו כמה כבר נחסך, כמה נשאר וכמה צריך להפקיד כל חודש."),
    ("💳 שלב 6 – חובות", "הזינו סכום מקורי, יתרה בתחילת השנה וריבית. ההחזרים נמשכים מהתנועות, "
     "ותראו כמה נפרע, כמה נשאר ומתי תהיו חופשיים מחובות."),
]
rr = 4
for head, body in steps:
    style(ws.cell(rr, 2, head), f=font(11, True, DARK), bg="EEF2F6", al=Alignment(horizontal="right", vertical="center", wrap_text=True), border=BOX)
    style(ws.cell(rr, 3, body), f=font(10), al=Alignment(horizontal="right", vertical="center", wrap_text=True), border=BOX)
    ws.row_dimensions[rr].height = 48
    rr += 1
rr += 1
style(ws.cell(rr, 2, "🎨 מקרא צבעים"), f=font(12, True, DARK), al=RIGHT)
rr += 1
legend = [(INPUT, "תא צהוב = תא למילוי (כל השאר מחושב אוטומטית – לא לגעת)")]
legend += [(COLORS[t][0], t) for t in TYPES]
legend += [("FDE2E2", "אדום = חריגה מהתקציב")]
for color, text in legend:
    style(ws.cell(rr, 2), bg=color, border=BOX)
    style(ws.cell(rr, 3, text), f=font(10), al=RIGHT, border=BOX)
    rr += 1
rr += 1
style(ws.cell(rr, 2, "⚠️ חשוב לדעת"), f=font(12, True, DARK), al=RIGHT)
tips = [
    "בלשונית 'תנועות' יש תנועות לדוגמה בחודש ינואר כדי שתראו איך הכל עובד. כשאתם מוכנים – סמנו אותן ומחקו (מקש Delete).",
    "שמות הקטגוריות בהגדרות צריכים להיות ייחודיים. שינוי שם בהגדרות מעדכן את כל הלשוניות; תנועות ישנות עם השם הקודם לא ייספרו.",
    "שמות תוכניות החיסכון והחובות נלקחים מההגדרות – הפקדה לחיסכון או החזר חוב פשוט נרשמים כתנועה עם אותה קטגוריה.",
    "הסכומים תמיד חיוביים – הסוג (הכנסה / הוצאה) נקבע לפי הקטגוריה.",
    "לשנה חדשה: שכפלו את הקובץ, שנו את השנה בהגדרות ומחקו את התנועות.",
]
for tip in tips:
    rr += 1
    ws.merge_cells(f"B{rr}:C{rr}")
    style(ws.cell(rr, 2, "• " + tip), f=font(10), al=Alignment(horizontal="right", vertical="center", wrap_text=True))
    ws.row_dimensions[rr].height = 30

wb.active = 0
wb.save(OUT)
print(f"saved {OUT}")
