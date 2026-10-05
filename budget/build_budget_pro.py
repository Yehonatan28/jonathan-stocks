"""בונה את "תמונת מצב פיננסית" – מערכת תקציב פשוטה: מזינים פעם בחודש, והכל מתעדכן לבד.

הרצה:
    python budget/build_budget_pro.py           -> budget/תקציב_תמונת_מצב.xlsx       (נקי)
    DEMO=1 python budget/build_budget_pro.py    -> budget/תקציב_תמונת_מצב_דוגמה.xlsx (עם 9 חודשי דוגמה)

עובד ב-Google Sheets וב-Excel (פונקציות סטנדרטיות בלבד, בלי פונקציות ייחודיות לאחד מהם).
"""
from __future__ import annotations

import os
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.utils import get_column_letter as L
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation

DEMO = os.environ.get("DEMO") == "1"
HERE = Path(__file__).parent
OUT = HERE / ("תקציב_תמונת_מצב_דוגמה.xlsx" if DEMO else "תקציב_תמונת_מצב.xlsx")
YEAR = 2026

# ------------------------------------------------------------------ design tokens
FONT = "Arial"
INK = "0F172A"        # טקסט ראשי
INK_2 = "334155"
MUTED = "64748B"
LINE = "E2E8F0"
CANVAS = "F1F5F9"     # רקע הדף
WHITE = "FFFFFF"
NAVY = "1E293B"
NAVY_2 = "334155"
INPUT_BG = "FFFBEB"   # תא להזנה
INPUT_BORDER = "FCD34D"
GOOD, BAD = "059669", "DC2626"

CUR = '₪#,##0;[Red]-₪#,##0;"–"'
CUR_PLAIN = '₪#,##0;-₪#,##0;"–"'
PCT = '0%;[Red]-0%;"–"'

# סוג: (כותרת, צבע ראשי, צבע בהיר)
TYPES = {
    "inc": ("💰 הכנסות", "2563EB", "DBEAFE"),
    "sav": ("🐷 חיסכון", "059669", "D1FAE5"),
    "bill": ("🧾 חשבונות ומנויים", "7C3AED", "EDE9FE"),
    "var": ("🛒 הוצאות משתנות", "EA580C", "FFEDD5"),
    "debt": ("💳 חובות", "DC2626", "FEE2E2"),
}
ORDER = ["inc", "sav", "bill", "var", "debt"]
EXPENSE = ["bill", "var", "debt"]
SHORT = {"inc": "הכנסות", "sav": "חיסכון", "bill": "חשבונות ומנויים", "var": "הוצאות משתנות", "debt": "חובות"}

# קטגוריות: (שם, תקציב חודשי, (ממוצע, פיזור) לנתוני דוגמה)
CATS = {
    "inc": [("משכורת 1", 14000, (14000, 0)), ("משכורת 2", 9500, (9500, 0)),
            ("עבודה נוספת / בונוסים", 1000, (1100, 900)), ("קצבאות", 320, (320, 0)),
            ("הכנסות אחרות", 0, (150, 150)), ("", None, None)],
    "sav": [("קרן חירום", 1500, (1500, 0)), ("חופשה משפחתית", 600, (600, 0)),
            ("רכב חדש", 800, (750, 150)), ("דירה", 1500, (1500, 0)),
            ("השקעות", 1000, (1000, 300))],
    "bill": [("שכר דירה / משכנתא", 5200, (5200, 0)), ("ארנונה", 450, (450, 0)),
             ("חשמל", 380, (380, 120)), ("מים וגז", 180, (180, 40)),
             ("ועד בית", 250, (250, 0)), ("תקשורת (אינטרנט + סלולר)", 300, (300, 0)),
             ("ביטוחים", 850, (850, 0)), ("מנויים וסטרימינג", 140, (140, 0)),
             ("חינוך וחוגים", 900, (920, 80)), ("", None, None)],
    "var": [("סופרמרקט", 2800, (2900, 350)), ("דלק ותחבורה", 750, (760, 150)),
            ("מסעדות ובילויים", 900, (950, 300)), ("ביגוד והנעלה", 400, (380, 300)),
            ("בריאות ותרופות", 200, (180, 120)), ("טיפוח", 200, (190, 80)),
            ("מתנות ואירועים", 300, (330, 250)), ("בית ותחזוקה", 300, (280, 220)),
            ("חופשות ופנאי", 500, (450, 500)), ("שונות", 300, (310, 150)), ("", None, None)],
    "debt": [("הלוואת רכב", 1100, (1100, 0)), ("הלוואה בנקאית", 800, (800, 0)),
             ("כרטיס אשראי – פריסה", 400, (400, 0)), ("", None, None)],
}
SAVINGS_GOALS = [(30000, 18000, "2027-06-30"), (12000, 2500, "2027-07-31"),
                 (40000, 15000, "2028-12-31"), (150000, 62000, "2030-12-31"),
                 (50000, 21000, "2028-12-31")]
DEBT_SETUP = [(38000, 0.055), (22000, 0.07), (4800, 0.12), (None, None)]

MONTHS = ["ינואר", "פברואר", "מרץ", "אפריל", "מאי", "יוני",
          "יולי", "אוגוסט", "ספטמבר", "אוקטובר", "נובמבר", "דצמבר"]
DEMO_MONTHS = 9

INP = "📥 הזנה חודשית"
DASH = "📊 תמונת מצב"
TRD = "📈 מגמות"
GOAL = "🎯 יעדים"
GUIDE = "📖 מדריך"


# ------------------------------------------------------------------ helpers
def q(name):
    return f"'{name}'"


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def side(color=LINE, style="thin"):
    return Side(style=style, color=color)


def al(h="center", v="center", wrap=False, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)


def put(ws, ref, value=None, *, f=None, bg=None, fmt=None, a=None, b=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    if f:
        c.font = f
    if bg:
        c.fill = fill(bg)
    if fmt:
        c.number_format = fmt
    if a:
        c.alignment = a
    if b:
        c.border = b
    return c


def area(ws, r1, r2, c1, c2, bg):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).fill = fill(bg)


def merge(ws, rng, value=None, **kw):
    ws.merge_cells(rng)
    first = rng.split(":")[0]
    return put(ws, first, value, **kw)


def setup(ws, tab, widths, *, landscape=True):
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.sheet_properties.tabColor = tab
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def banner(ws, row, c1, c2, title, subtitle):
    rng = f"{L(c1)}{row}:{L(c2)}{row}"
    area(ws, row, row + 1, c1, c2, NAVY)
    merge(ws, rng, title, f=font(22, True, WHITE), a=al("right", indent=1))
    merge(ws, f"{L(c1)}{row + 1}:{L(c2)}{row + 1}", subtitle, f=font(10, False, "CBD5E1"), a=al("right", indent=1))
    ws.row_dimensions[row].height = 44
    ws.row_dimensions[row + 1].height = 22


def section_title(ws, ref, text, color=INK):
    put(ws, ref, text, f=font(13, True, color), a=al("right"))
    ws.row_dimensions[ws[ref].row].height = 26


def header(ws, row, c1, labels, bg=NAVY, color=WHITE):
    for i, lab in enumerate(labels):
        c = ws.cell(row, c1 + i, lab)
        c.font = font(10, True, color)
        c.fill = fill(bg)
        c.alignment = al(wrap=True)
        c.border = Border(bottom=side(bg))
    ws.row_dimensions[row].height = 30


def bar_formula(ratio_ref, n=10):
    r = f"MAX(0,MIN(1,{ratio_ref}))"
    return f'REPT("█",ROUND({r}*{n},0))&REPT("░",{n}-ROUND({r}*{n},0))'


def card(ws, r, c1, c2, label, value, sub, accent, fmt=CUR_PLAIN):
    """כרטיס KPI: פס צבע עליון, תווית, מספר גדול ושורת הסבר. הרווח בין כרטיסים = גבול בצבע הרקע."""
    gap = side(CANVAS, "thick")
    for rr in range(r, r + 3):
        for cc in range(c1, c2 + 1):
            cell = ws.cell(rr, cc)
            cell.fill = fill(WHITE)
            cell.border = Border(
                left=gap if cc == c2 else None,  # RTL: עמודה אחרונה = צד שמאל
                right=gap if cc == c1 else None,
                top=side(accent, "thick") if rr == r else None,
                bottom=side(LINE) if rr == r + 2 else None,
            )
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    ws.merge_cells(start_row=r + 1, start_column=c1, end_row=r + 1, end_column=c2)
    ws.merge_cells(start_row=r + 2, start_column=c1, end_row=r + 2, end_column=c2)
    a = ws.cell(r, c1, label)
    a.font, a.alignment = font(10, True, MUTED), al("right", "bottom", indent=1)
    b = ws.cell(r + 1, c1, value)
    b.font, b.alignment, b.number_format = font(22, True, INK), al("right", indent=1), fmt
    s = ws.cell(r + 2, c1, sub)
    s.font, s.alignment = font(9, False, MUTED), al("right", "top", indent=1)
    ws.row_dimensions[r].height = 24
    ws.row_dimensions[r + 1].height = 36
    ws.row_dimensions[r + 2].height = 22


def money(x):
    return f'"₪"&TEXT({x},"#,##0")'


def pct_labels():
    return DataLabelList(showPercent=True, showVal=False, showCatName=False,
                         showSerName=False, showLegendKey=False, showLeaderLines=False)


def place(ws, ch, rng):
    """מצמיד גרף לטווח תאים (עוגן דו-תאי) – כך הוא לא זז בתצוגת ימין-לשמאל."""
    c1, r1, c2, r2 = range_boundaries(rng)
    anchor = TwoCellAnchor()
    anchor._from = AnchorMarker(col=c1 - 1, row=r1 - 1)
    anchor.to = AnchorMarker(col=c2, row=r2)
    ch.anchor = anchor
    ws.add_chart(ch)


def style_chart(ch, title, w, h):
    ch.title = title
    ch.width, ch.height = w, h
    ch.style = 2
    if getattr(ch, "legend", None) is not None:
        ch.legend.position = "b"
    return ch


# ------------------------------------------------------------------ workbook
wb = Workbook()
wb.remove(wb.active)
ws_dash = wb.create_sheet(DASH)
ws_in = wb.create_sheet(INP)
ws_tr = wb.create_sheet(TRD)
ws_goal = wb.create_sheet(GOAL)
ws_guide = wb.create_sheet(GUIDE)

# ================================================================== 📥 הזנה חודשית
ws = ws_in
MC1, MC12 = 3, 14            # עמודות C..N = ינואר..דצמבר
COL_TOT, COL_AVG = 15, 16    # O, P
setup(ws, "F59E0B", {"A": 28, "B": 13, **{L(c): 11.5 for c in range(MC1, MC12 + 1)}, "O": 13, "P": 13})
banner(ws, 1, 1, 16, "📥 הזנה חודשית",
       "פעם בחודש: ממלאים את העמודה של החודש עם הסכומים מדף הבנק ומפירוט האשראי. כל השאר מתעדכן לבד ✨")
put(ws, "A3", "שנת התקציב", f=font(10, True, INK_2), a=al("right", indent=1))
put(ws, "B3", YEAR, f=font(12, True, INK), bg=INPUT_BG, a=al(),
    b=Border(*(side(INPUT_BORDER),) * 4))
merge(ws, "D3:P3", "💡 תאים בצבע שמנת = להזנה. אפשר לשנות שמות קטגוריות ולמלא שורות ריקות – הכל מתעדכן אוטומטית.",
      f=font(9, False, MUTED), a=al("right"))
ws.row_dimensions[3].height = 26
YEAR_REF = f"{q(INP)}!$B$3"
# שורת עזר מוסתרת: מספרי החודשים
for i in range(12):
    ws.cell(4, MC1 + i, i + 1)
ws.row_dimensions[4].hidden = True
header(ws, 5, 1, ["קטגוריה", "תקציב חודשי"] + MONTHS + ['סה"כ שנתי', "ממוצע לחודש"])

dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        showErrorMessage=True, errorTitle="סכום לא תקין", error="הזן סכום חיובי (בלי סימן ₪)")
ws.add_data_validation(dv_num)

row = 6
SEC = {}       # type -> header row
CAT_ROWS = {}  # type -> list of rows
for t in ORDER:
    title_txt, accent, light = TYPES[t]
    SEC[t] = row
    rows = list(range(row + 1, row + 1 + len(CATS[t])))
    CAT_ROWS[t] = rows
    # שורת כותרת = סיכום הסעיף
    c = ws.cell(row, 1, title_txt)
    c.font, c.alignment = font(11, True, accent), al("right", indent=1)
    for cc in range(1, COL_AVG + 1):
        cell = ws.cell(row, cc)
        cell.fill = fill(light)
        cell.border = Border(top=side(accent, "medium"))
        if cc >= 2:
            col = L(cc)
            cell.value = f"=SUM({col}{rows[0]}:{col}{rows[-1]})"
            cell.font, cell.alignment, cell.number_format = font(10, True, accent), al(), CUR_PLAIN
    ws.row_dimensions[row].height = 26
    for i, r_ in enumerate(rows):
        name, budget, _ = CATS[t][i]
        put(ws, f"A{r_}", name or None, f=font(10, False, INK), bg=INPUT_BG, a=al("right", indent=1),
            b=Border(bottom=side(), right=side(accent, "thick")))
        put(ws, f"B{r_}", budget, f=font(10, False, INK_2), bg=INPUT_BG, fmt=CUR_PLAIN, a=al(),
            b=Border(bottom=side(), left=side(LINE)))
        for cc in range(MC1, MC12 + 1):
            put(ws, f"{L(cc)}{r_}", None, f=font(10), bg=WHITE, fmt=CUR_PLAIN, a=al(), b=Border(bottom=side(), left=side("F1F5F9")))
        dv_num.add(f"B{r_}:{L(MC12)}{r_}")
        put(ws, f"O{r_}", f"=SUM(C{r_}:N{r_})", f=font(10, True, INK), bg="F8FAFC", fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
        put(ws, f"P{r_}", f'=IF(COUNT(C{r_}:N{r_})=0,0,AVERAGE(C{r_}:N{r_}))', f=font(10, False, MUTED), bg="F8FAFC",
            fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
        ws.row_dimensions[r_].height = 21
        # חריגה מהתקציב בהוצאות – אדום עדין
        if t in EXPENSE:
            ws.conditional_formatting.add(
                f"C{r_}:N{r_}",
                FormulaRule(formula=[f"AND(ISNUMBER(C{r_}),$B{r_}>0,C{r_}>$B{r_})"],
                            font=Font(name=FONT, color=BAD, bold=True), fill=fill("FEF2F2")))
    row = rows[-1] + 2

# ---- סיכום חודשי
SUM0 = row
c = ws.cell(SUM0, 1, "📊 סיכום")
c.font, c.alignment = font(11, True, WHITE), al("right", indent=1)
area(ws, SUM0, SUM0, 1, COL_AVG, NAVY)
ws.row_dimensions[SUM0].height = 26
R_INC, R_EXP, R_SAV, R_BAL, R_RATE, R_FILL = range(SUM0 + 1, SUM0 + 7)
summary = [
    (R_INC, "סה\"כ הכנסות", lambda col: f"={col}{SEC['inc']}", CUR_PLAIN),
    (R_EXP, "סה\"כ הוצאות (חשבונות + משתנות + חובות)",
     lambda col: "=" + "+".join(f"{col}{SEC[t]}" for t in EXPENSE), CUR_PLAIN),
    (R_SAV, "סה\"כ הופקד לחיסכון", lambda col: f"={col}{SEC['sav']}", CUR_PLAIN),
    (R_BAL, "נשאר בעו\"ש", lambda col: f"={col}{R_INC}-{col}{R_EXP}-{col}{R_SAV}", CUR),
    (R_RATE, "שיעור חיסכון", lambda col: f'=IF(N({col}{R_INC})=0,"",{col}{R_SAV}/{col}{R_INC})', "0%"),
]
for r_, label, fn, fmt in summary:
    put(ws, f"A{r_}", label, f=font(10, True, INK), bg=WHITE, a=al("right", indent=1), b=Border(bottom=side()))
    for cc in range(2, COL_AVG + 1):
        col = L(cc)
        if cc == COL_AVG and r_ != R_RATE:
            formula = f'=IF(MAX($C${R_FILL}:$N${R_FILL})=0,0,O{r_}/MAX($C${R_FILL}:$N${R_FILL}))'
        else:
            formula = fn(col)
        put(ws, f"{col}{r_}", formula, f=font(10, r_ in (R_BAL,), INK), bg=WHITE if cc <= MC12 else "F8FAFC",
            fmt=fmt, a=al(), b=Border(bottom=side()))
    ws.row_dimensions[r_].height = 21
ws.cell(R_BAL, 1).font = font(10, True, NAVY)
for cc in range(1, COL_AVG + 1):
    ws.cell(R_BAL, cc).fill = fill("F8FAFC")
# שורת סטטוס מילוי (✅) – משמשת לזיהוי החודש האחרון שמולא
put(ws, f"A{R_FILL}", "מולא?", f=font(10, True, MUTED), a=al("right", indent=1))
all_cat_rows = [r_ for t in ORDER for r_ in CAT_ROWS[t]]
for cc in range(MC1, MC12 + 1):
    col = L(cc)
    counts = "+".join(f"COUNT({col}{CAT_ROWS[t][0]}:{col}{CAT_ROWS[t][-1]})" for t in ORDER)
    put(ws, f"{col}{R_FILL}", f"=IF({counts}>0,{col}$4,0)", f=font(11, True, GOOD), fmt='"✅";;"—"', a=al())
ws.cell(R_FILL, 2).value = None
ws.cell(R_RATE, 2).value = f'=IF(N(B{R_INC})=0,"",B{R_SAV}/B{R_INC})'
ws.cell(R_RATE, COL_TOT).value = f'=IF(N(O{R_INC})=0,"",O{R_SAV}/O{R_INC})'
ws.cell(R_RATE, COL_AVG).value = f'=O{R_RATE}'
ws.freeze_panes = "C6"

# נתוני דוגמה
if DEMO:
    rnd = random.Random(7)
    for t in ORDER:
        for i, r_ in enumerate(CAT_ROWS[t]):
            spec = CATS[t][i][2]
            if not spec:
                continue
            mean, spread = spec
            for m in range(DEMO_MONTHS):
                v = mean + (rnd.uniform(-spread, spread) if spread else 0)
                if t == "var" and m in (3, 6) and CATS[t][i][0] in ("חופשות ופנאי", "מתנות ואירועים"):
                    v *= 2.6  # חגים / חופש גדול
                if t == "inc" and CATS[t][i][0].startswith("עבודה") and m == 3:
                    v += 2500
                v = max(0, round(v / 10) * 10)
                ws.cell(r_, MC1 + m, v)


def in_ref(r_, col=None):
    return f"{q(INP)}!${col}${r_}" if col else f"{q(INP)}!$C${r_}:$N${r_}"


MONTH_NUMS = f"{q(INP)}!$C$4:$N$4"
FILL_ROW = f"{q(INP)}!$C${R_FILL}:$N${R_FILL}"

# ================================================================== 📊 תמונת מצב
ws = ws_dash
setup(ws, NAVY, {"A": 2, **{L(c): 12.5 for c in range(2, 14)}, "N": 2})
area(ws, 1, 75, 1, 14, CANVAS)
banner(ws, 2, 2, 13, "📊 תמונת מצב פיננסית", None)
ws.row_dimensions[1].height = 10

# תאי עזר (עמודות מוסתרות X..AD)
HX = 24  # X
for i, m in enumerate(MONTHS):
    ws.cell(3 + i, HX, m)
SEL = "$Y$1"
LAST = "$Y$2"
MNAME = "$Y$3"
ws["Y2"] = f"=MAX({FILL_ROW})"
AUTO = "אוטומטי – החודש האחרון שמולא"
ws["Y1"] = f'=IF(OR($C$5="",$C$5="{AUTO}"),MAX(1,Y2),IFERROR(MATCH($C$5,$X$3:$X$14,0),MAX(1,Y2)))'
ws["Y3"] = "=INDEX($X$3:$X$14,$Y$1)"
ws["Y4"] = '=IF($Y$1=1,"",INDEX($X$3:$X$14,$Y$1-1))'
for col in ("X", "Y", "Z", "AA", "AB", "AC", "AD"):
    ws.column_dimensions[col].hidden = True

ws["B3"] = (f'=IF({LAST}=0,"👋 ברוכים הבאים! מלאו את החודש הראשון בלשונית \'{INP}\' והלוח יתעורר לחיים",'
            f'"מציג את "&{MNAME}&" "&{YEAR_REF}&"  •  מולאו "&COUNTIF({FILL_ROW},">0")&" מתוך 12 חודשים")')

# בורר חודש
put(ws, "B5", "📅 חודש לתצוגה", f=font(10, True, INK_2), a=al("right", indent=1))
merge(ws, "C5:E5", AUTO, f=font(10, True, INK), bg=WHITE, a=al("right", indent=1))
for cc in range(3, 6):
    ws.cell(5, cc).border = Border(top=side(INPUT_BORDER), bottom=side(INPUT_BORDER),
                                   left=side(INPUT_BORDER) if cc == 5 else None,
                                   right=side(INPUT_BORDER) if cc == 3 else None)
    ws.cell(5, cc).fill = fill(INPUT_BG)
dv_m = DataValidation(type="list", formula1='"' + ",".join([AUTO] + MONTHS) + '"', allow_blank=True)
ws.add_data_validation(dv_m)
dv_m.add("C5")
merge(ws, "F5:M5", "← לחצו כדי לבחור חודש. במצב 'אוטומטי' הלוח מציג תמיד את החודש האחרון שמילאתם.",
      f=font(9, False, MUTED), a=al("right"))
ws.row_dimensions[5].height = 26


def cur(r_):    # ערך החודש הנבחר
    return f"INDEX({in_ref(r_)},{SEL})"


def prev(r_):   # ערך החודש הקודם
    return f"IF({SEL}=1,0,INDEX({in_ref(r_)},{SEL}-1))"


def ytd(r_):    # מתחילת השנה עד החודש הנבחר
    return f"SUMPRODUCT(({MONTH_NUMS}<={SEL})*{in_ref(r_)})"


def bud(r_):
    return in_ref(r_, "B")


inc_b, exp_b, sav_b = bud(R_INC), bud(R_EXP), bud(R_SAV)
card(ws, 7, 2, 4, "💰 הכנסות החודש", f"={cur(R_INC)}",
     f'=IF({inc_b}=0,"",IF({cur(R_INC)}>={inc_b},"✅ ","⏳ ")&"תכנון: "&{money(inc_b)})', TYPES["inc"][1])
card(ws, 7, 5, 7, "🧾 הוצאות החודש", f"={cur(R_EXP)}",
     f'=IF({exp_b}=0,"",IF({cur(R_EXP)}<={exp_b},"✅ "&{money(f"{exp_b}-{cur(R_EXP)}")}&" מתחת לתקציב",'
     f'"⚠️ "&{money(f"{cur(R_EXP)}-{exp_b}")}&" מעל התקציב"))', TYPES["var"][1])
card(ws, 7, 8, 10, "🐷 הופקד לחיסכון", f"={cur(R_SAV)}",
     f'=IF({sav_b}=0,"",IF({cur(R_SAV)}>={sav_b},"✅ עמדת ביעד ","⏳ יעד: ")&{money(sav_b)})', TYPES["sav"][1])
card(ws, 7, 11, 13, "🏦 נשאר בעו\"ש", f"={cur(R_BAL)}",
     f'=IF({SEL}=1,"החודש הראשון בשנה",IF({cur(R_BAL)}>={prev(R_BAL)},"▲ ","▼ ")&{money(f"ABS({cur(R_BAL)}-{prev(R_BAL)})")}&" לעומת "&$Y$4)',
     NAVY)

ws["Z10"] = f"={ytd(R_INC)}"
ws["Z11"] = f"={ytd(R_EXP)}"
ws["Z12"] = f"={ytd(R_SAV)}"
ws["Z13"] = f"={ytd(R_BAL)}"
card(ws, 11, 2, 4, "📅 הכנסות מתחילת השנה", "=$Z$10", f'="ממוצע לחודש: "&{money("$Z$10/" + SEL)}', TYPES["inc"][1])
card(ws, 11, 5, 7, "📅 הוצאות מתחילת השנה", "=$Z$11", f'="ממוצע לחודש: "&{money("$Z$11/" + SEL)}', TYPES["var"][1])
card(ws, 11, 8, 10, "📅 נחסך מתחילת השנה", "=$Z$12",
     f'="בקצב הזה: כ-"&{money("$Z$12/" + SEL + "*12")}&" עד סוף השנה"', TYPES["sav"][1])
card(ws, 11, 11, 13, "🎯 שיעור חיסכון שנתי", '=IF($Z$10=0,0,$Z$12/$Z$10)',
     '=IF($Z$10=0,"",IF($Z$12/$Z$10>=0.2,"מצוין! מעל היעד המומלץ (20%) 🎉",IF($Z$12/$Z$10>=0.1,"טוב – היעד המומלץ הוא 20%","כדאי לשאוף ל-10%–20% מההכנסה")))',
     NAVY, fmt="0%")

# מצב ריק: לפני שמולא חודש ראשון – בלי הודעות מטעות
for rr_ in (9, 13):
    for cc_ in (2, 5, 8, 11):
        c_ = ws.cell(rr_, cc_)
        c_.value = '=IF($Y$2=0,"—",' + c_.value[1:] + ")"

# ---- תקציב מול בפועל
section_title(ws, "B15", "⚖️ תקציב מול בפועל – החודש")
header(ws, 16, 2, ["סעיף", "תקציב", "בפועל", "הפרש", "ניצול", "סטטוס"])
T0 = 17
for i, t in enumerate(ORDER):
    r_ = T0 + i
    accent, light = TYPES[t][1], TYPES[t][2]
    hr = SEC[t]
    put(ws, f"B{r_}", SHORT[t], f=font(10, True, accent), bg=WHITE, a=al("right", indent=1),
        b=Border(right=side(accent, "thick"), bottom=side()))
    put(ws, f"C{r_}", f"={bud(hr)}", f=font(10, False, INK_2), bg=WHITE, fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
    put(ws, f"D{r_}", f"={cur(hr)}", f=font(10, True, INK), bg=WHITE, fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
    diff = f"=D{r_}-C{r_}" if t in ("inc", "sav") else f"=C{r_}-D{r_}"
    put(ws, f"E{r_}", diff, f=font(10, False, INK), bg=WHITE, fmt=CUR, a=al(), b=Border(bottom=side()))
    put(ws, f"F{r_}", f'=IF(C{r_}=0,"",{bar_formula(f"D{r_}/C{r_}", 8)})', f=font(9, False, accent), bg=WHITE,
        a=al(), b=Border(bottom=side()))
    if t in ("inc", "sav"):
        status = f'=IF(C{r_}=0,"",IF(D{r_}>=C{r_},"✅ הושג",TEXT(D{r_}/C{r_},"0%")&" מהיעד"))'
    else:
        status = f'=IF(C{r_}=0,"",IF(D{r_}<=C{r_},"✅ בתקציב","🔴 חריגה"))'
    put(ws, f"G{r_}", status, f=font(10, True, INK_2), bg=WHITE, a=al(), b=Border(bottom=side()))
    ws.row_dimensions[r_].height = 24
    if t in EXPENSE:
        ws.conditional_formatting.add(f"G{r_}", FormulaRule(formula=[f"D{r_}>C{r_}"], font=Font(name=FONT, color=BAD, bold=True)))
        ws.conditional_formatting.add(f"E{r_}", FormulaRule(formula=[f"E{r_}<0"], fill=fill("FEF2F2")))
for i in range(5):
    for col_ in ("E", "G"):
        c_ = ws[f"{col_}{T0 + i}"]
        c_.value = '=IF($Y$2=0,"",' + c_.value[1:] + ")"
TB = T0 + 5
put(ws, f"B{TB}", "🏦 נשאר בעו\"ש", f=font(10, True, WHITE), bg=NAVY, a=al("right", indent=1))
put(ws, f"C{TB}", f"={bud(R_BAL)}", f=font(10, True, WHITE), bg=NAVY, fmt=CUR, a=al())
put(ws, f"D{TB}", f"={cur(R_BAL)}", f=font(10, True, WHITE), bg=NAVY, fmt=CUR, a=al())
put(ws, f"E{TB}", f"=D{TB}-C{TB}", f=font(10, True, WHITE), bg=NAVY, fmt=CUR_PLAIN, a=al())
for col in ("F", "G"):
    put(ws, f"{col}{TB}", None, bg=NAVY)
ws.row_dimensions[TB].height = 26

# ---- עזר: כל קטגוריות ההוצאה עבור החודש הנבחר (עמודות מוסתרות Z..AC)
exp_rows = [r_ for t in EXPENSE for r_ in CAT_ROWS[t]]
H0 = 20
for i, r_ in enumerate(exp_rows):
    hr = H0 + i
    ws[f"Z{hr}"] = f'={in_ref(r_, "A")}&""'
    ws[f"AA{hr}"] = f"=N({bud(r_)})"
    ws[f"AB{hr}"] = f"=N({cur(r_)})"
    ws[f"AC{hr}"] = f'=IF(AND(Z{hr}<>"",AB{hr}>0),AB{hr}+ROW()/1000000,0)'
H1 = H0 + len(exp_rows) - 1
HZ, HA, HB, HC = (f"${c}${H0}:${c}${H1}" for c in ("Z", "AA", "AB", "AC"))

# ---- 5 ההוצאות הגדולות
section_title(ws, "B25", "🏆 5 ההוצאות הגדולות של החודש")
header(ws, 26, 2, ["#", "קטגוריה", "", "סכום", "% מההוצאות", "מול תקציב"])
ws.merge_cells("C26:D26")
for k in range(1, 6):
    r_ = 26 + k
    key = f"LARGE({HC},{k})"
    pos = f"MATCH({key},{HC},0)"
    put(ws, f"B{r_}", ["🥇", "🥈", "🥉", "4", "5"][k - 1], f=font(12, True, INK_2), bg=WHITE, a=al(), b=Border(bottom=side()))
    merge(ws, f"C{r_}:D{r_}", f'=IF({key}<=0,"",INDEX({HZ},{pos}))', f=font(10, True, INK), bg=WHITE, a=al("right", indent=1))
    ws[f"D{r_}"].border = ws[f"C{r_}"].border = Border(bottom=side())
    ws[f"D{r_}"].fill = fill(WHITE)
    put(ws, f"E{r_}", f'=IF({key}<=0,"",INDEX({HB},{pos}))', f=font(10, True, INK), bg=WHITE, fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
    put(ws, f"F{r_}", f'=IF(OR(E{r_}="",N(D{T0 + 2})+N(D{T0 + 3})+N(D{T0 + 4})=0),"",E{r_}/(D{T0 + 2}+D{T0 + 3}+D{T0 + 4}))',
        f=font(10, False, INK_2), bg=WHITE, fmt="0%", a=al(), b=Border(bottom=side()))
    put(ws, f"G{r_}", f'=IF(E{r_}="","",IF(INDEX({HA},{pos})=0,"—",IF(E{r_}<=INDEX({HA},{pos}),"✅ בתקציב","🔴 +"&TEXT(E{r_}-INDEX({HA},{pos}),"#,##0"))))',
        f=font(10, False, INK_2), bg=WHITE, a=al(), b=Border(bottom=side()))
    ws.row_dimensions[r_].height = 24

# ---- תובנות
section_title(ws, "H15", "💡 תובנות אוטומטיות")
merge(ws, "H16:M16", "מה המספרים אומרים החודש", f=font(10, True, WHITE), bg=NAVY, a=al("right", indent=1))
area(ws, 16, 16, 8, 13, NAVY)
inc_c, exp_c, sav_c = cur(R_INC), cur(R_EXP), cur(R_SAV)
over_n = f"SUMPRODUCT(({HB}>{HA})*({HA}>0))"
insights = [
    f'=IF({LAST}=0,"👋 עוד לא הוזנו נתונים – פתחו את \'{INP}\' ומלאו את העמודה של החודש.",'
    f'IF({inc_c}=0,"ℹ️ לא הוזנו הכנסות בחודש הזה.","🐷 החודש חסכתם "&TEXT({sav_c}/{inc_c},"0%")&" מההכנסה"&'
    f'IF({sav_c}/{inc_c}>=0.2," – מצוין! 🎉",IF({sav_c}/{inc_c}>=0.1," – יפה, היעד המומלץ הוא 20% 👍"," – נסו להגדיל לפחות ל-10%"))))',
    f'=IF({LAST}=0,"",IF({over_n}=0,"✅ כל קטגוריות ההוצאה עמדו בתקציב החודש – כל הכבוד!",'
    f'"⚠️ "&{over_n}&" קטגוריות חרגו מהתקציב החודש (מסומנות באדום בלשונית ההזנה)"))',
    f'=IF(C27="","","💸 ההוצאה הגדולה ביותר: "&C27&" – "&{money("E27")}&" ("&TEXT(F27,"0%")&" מכלל ההוצאות)")',
    f'=IF(OR({SEL}=1,{LAST}=0,{prev(R_EXP)}=0),"",IF({exp_c}>{prev(R_EXP)},"📈 ההוצאות עלו ב-","📉 ההוצאות ירדו ב-")&'
    f'TEXT(ABS({exp_c}/{prev(R_EXP)}-1),"0%")&" לעומת "&$Y$4)',
    f'=IF({LAST}=0,"","🔮 תחזית לסוף השנה (לפי הקצב עד עכשיו): הכנסות "&{money("$Z$10/" + SEL + "*12")}&'
    f'" • חיסכון "&{money("$Z$12/" + SEL + "*12")})',
]
for i, fml in enumerate(insights):
    r_ = T0 + i
    merge(ws, f"H{r_}:M{r_}", fml, f=font(10, False, INK), bg=WHITE, a=al("right", indent=1, wrap=True))
    for cc in range(8, 14):
        ws.cell(r_, cc).fill = fill(WHITE)
        ws.cell(r_, cc).border = Border(bottom=side())
    ws.row_dimensions[r_].height = 26

# ---- גרפים בלוח
donut = DoughnutChart()
donut.add_data(Reference(ws, min_col=4, min_row=T0 + 1, max_row=T0 + 4), titles_from_data=False)
donut.set_categories(Reference(ws, min_col=2, min_row=T0 + 1, max_row=T0 + 4))
for i, t in enumerate(["sav", "bill", "var", "debt"]):
    pt = DataPoint(idx=i)
    pt.graphicalProperties.solidFill = TYPES[t][1]
    pt.graphicalProperties.line.solidFill = WHITE
    donut.series[0].dPt.append(pt)
donut.dataLabels = pct_labels()
donut.holeSize = 58
style_chart(donut, "לאן הלך הכסף החודש", 13.2, 7.4)
place(ws, donut, "I24:M38")

bar = BarChart()
bar.type = "col"
bar.add_data(Reference(ws, min_col=3, max_col=4, min_row=16, max_row=T0 + 4), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=2, min_row=T0, max_row=T0 + 4))
bar.series[0].graphicalProperties.solidFill = "CBD5E1"
bar.series[0].graphicalProperties.line.noFill = True
bar.series[1].graphicalProperties.solidFill = "2563EB"
bar.series[1].graphicalProperties.line.noFill = True
bar.gapWidth = 60
bar.y_axis.numFmt = "#,##0"
bar.y_axis.majorGridlines = None
bar.x_axis.scaling.orientation = "maxMin"
style_chart(bar, "תקציב מול בפועל", 13.2, 7.4)
place(ws, bar, "B40:G55")

trend = BarChart()
trend.type = "col"
for cc, color in ((3, TYPES["inc"][1]), (7, TYPES["var"][1])):
    trend.add_data(Reference(ws_tr, min_col=cc, min_row=6, max_row=18), titles_from_data=True)
    trend.series[-1].graphicalProperties.solidFill = color
    trend.series[-1].graphicalProperties.line.noFill = True
trend.set_categories(Reference(ws_tr, min_col=2, min_row=7, max_row=18))
trend.gapWidth = 60
trend.y_axis.numFmt = "#,##0"
trend.y_axis.majorGridlines = None
trend.x_axis.scaling.orientation = "maxMin"
style_chart(trend, "הכנסות מול הוצאות – כל השנה", 13.2, 7.4)
place(ws, trend, "I40:M55")

ws.freeze_panes = "A4"

# ================================================================== 📈 מגמות
ws = ws_tr
setup(ws, "2563EB", {"A": 2, "B": 24, **{L(c): 12 for c in range(3, 16)}, "P": 2})
area(ws, 1, 110, 1, 16, CANVAS)
banner(ws, 2, 2, 15, "📈 מגמות לאורך השנה", "איך הכסף זז מחודש לחודש – הכנסות, הוצאות, חיסכון ומפת חום של כל קטגוריה")
ws.row_dimensions[1].height = 10
section_title(ws, "B5", "📅 סיכום לפי חודש")
TR_H = 6
header(ws, TR_H, 2, ["חודש", "הכנסות", "חשבונות ומנויים", "הוצאות משתנות", "חובות", 'סה"כ הוצאות',
                     "חיסכון", "נשאר בעו\"ש", "שיעור חיסכון", "חיסכון מצטבר"])
for mi in range(12):
    r_ = TR_H + 1 + mi
    col = L(MC1 + mi)
    filled = f"{q(INP)}!{col}${R_FILL}=0"
    src = {3: f"{col}{R_INC}", 4: f"{col}{SEC['bill']}", 5: f"{col}{SEC['var']}", 6: f"{col}{SEC['debt']}",
           7: f"{col}{R_EXP}", 8: f"{col}{R_SAV}", 9: f"{col}{R_BAL}"}
    band = WHITE if mi % 2 == 0 else "F8FAFC"
    put(ws, f"B{r_}", MONTHS[mi], f=font(10, True, INK), bg=band, a=al("right", indent=1), b=Border(bottom=side()))
    for cc, ref in src.items():
        put(ws, f"{L(cc)}{r_}", f'=IF({filled},"",{q(INP)}!{ref})', f=font(10, cc in (7, 9), INK),
            bg=band, fmt=CUR if cc == 9 else CUR_PLAIN, a=al(), b=Border(bottom=side()))
    put(ws, f"J{r_}", f'=IF(OR(C{r_}="",N(C{r_})=0),"",H{r_}/C{r_})', f=font(10, False, INK_2), bg=band, fmt="0%", a=al(), b=Border(bottom=side()))
    put(ws, f"K{r_}", f'=IF(H{r_}="","",SUM(H${TR_H + 1}:H{r_}))', f=font(10, True, TYPES["sav"][1]), bg=band, fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
    ws.row_dimensions[r_].height = 21
TR_T, TR_A = TR_H + 13, TR_H + 14
put(ws, f"B{TR_T}", 'סה"כ', f=font(10, True, WHITE), bg=NAVY, a=al("right", indent=1))
put(ws, f"B{TR_A}", "ממוצע לחודש", f=font(10, True, INK), bg="E2E8F0", a=al("right", indent=1))
for cc in range(3, 12):
    col = L(cc)
    if cc == 10:
        t_f = f'=IF(N(C{TR_T})=0,"",H{TR_T}/C{TR_T})'
        a_f = t_f.replace(str(TR_T), str(TR_A))
        fmt = "0%"
    elif cc == 11:
        t_f, a_f, fmt = f"=MAX(K{TR_H + 1}:K{TR_H + 12})", "", CUR_PLAIN
    else:
        t_f = f"=SUM({col}{TR_H + 1}:{col}{TR_H + 12})"
        a_f = f'=IF(COUNT({col}{TR_H + 1}:{col}{TR_H + 12})=0,"",AVERAGE({col}{TR_H + 1}:{col}{TR_H + 12}))'
        fmt = CUR
    put(ws, f"{col}{TR_T}", t_f, f=font(10, True, WHITE), bg=NAVY, fmt=fmt, a=al())
    put(ws, f"{col}{TR_A}", a_f or None, f=font(10, False, INK), bg="E2E8F0", fmt=fmt, a=al())
ws.row_dimensions[TR_T].height = ws.row_dimensions[TR_A].height = 24

cats_ref = Reference(ws, min_col=2, min_row=TR_H + 1, max_row=TR_H + 12)
c1 = BarChart()
c1.type = "col"
for cc, color in ((3, TYPES["inc"][1]), (7, TYPES["var"][1]), (8, TYPES["sav"][1])):
    c1.add_data(Reference(ws, min_col=cc, min_row=TR_H, max_row=TR_H + 12), titles_from_data=True)
    c1.series[-1].graphicalProperties.solidFill = color
    c1.series[-1].graphicalProperties.line.noFill = True
c1.set_categories(cats_ref)
c1.gapWidth = 50
c1.y_axis.numFmt = "#,##0"
c1.y_axis.majorGridlines = None
c1.x_axis.scaling.orientation = "maxMin"
style_chart(c1, "הכנסות, הוצאות וחיסכון לפי חודש", 17.5, 8)
place(ws, c1, "B22:G38")

c2 = BarChart()
c2.type = "col"
c2.grouping = "stacked"
c2.overlap = 100
for cc, t in ((4, "bill"), (5, "var"), (6, "debt")):
    c2.add_data(Reference(ws, min_col=cc, min_row=TR_H, max_row=TR_H + 12), titles_from_data=True)
    c2.series[-1].graphicalProperties.solidFill = TYPES[t][1]
    c2.series[-1].graphicalProperties.line.noFill = True
c2.set_categories(cats_ref)
c2.gapWidth = 50
c2.y_axis.numFmt = "#,##0"
c2.y_axis.majorGridlines = None
c2.x_axis.scaling.orientation = "maxMin"
style_chart(c2, "ממה מורכבות ההוצאות", 17.5, 8)
place(ws, c2, "I22:O38")

c3 = BarChart()
c3.type = "col"
c3.gapWidth = 50
c3.add_data(Reference(ws, min_col=11, min_row=TR_H, max_row=TR_H + 12), titles_from_data=True)
c3.set_categories(cats_ref)
c3.series[0].graphicalProperties.solidFill = TYPES["sav"][1]
c3.series[0].graphicalProperties.line.noFill = True
c3.y_axis.numFmt = "#,##0"
c3.y_axis.majorGridlines = None
c3.x_axis.scaling.orientation = "maxMin"
c3.legend = None
style_chart(c3, "חיסכון מצטבר מתחילת השנה", 17.5, 8)
place(ws, c3, "B40:G56")

c4 = BarChart()
c4.type = "col"
c4.gapWidth = 50
c4.add_data(Reference(ws, min_col=9, min_row=TR_H, max_row=TR_H + 12), titles_from_data=True)
c4.set_categories(cats_ref)
c4.series[0].graphicalProperties.solidFill = NAVY
c4.series[0].graphicalProperties.line.noFill = True
c4.y_axis.numFmt = "#,##0"
c4.y_axis.majorGridlines = None
c4.x_axis.scaling.orientation = "maxMin"
c4.legend = None
style_chart(c4, "נשאר בעו\"ש בכל חודש", 17.5, 8)
place(ws, c4, "I40:O56")

# ---- מפת חום
HM0 = 58
section_title(ws, f"B{HM0}", "🌡️ מפת חום – כמה הוצאתם בכל קטגוריה בכל חודש (כהה = יותר)")
header(ws, HM0 + 1, 2, ["קטגוריה"] + MONTHS + ['סה"כ'])
r_ = HM0 + 2
for t in EXPENSE:
    accent, light = TYPES[t][1], TYPES[t][2]
    put(ws, f"B{r_}", TYPES[t][0], f=font(10, True, accent), bg=light, a=al("right", indent=1))
    area(ws, r_, r_, 3, 15, light)
    r_ += 1
    first = r_
    for cr in CAT_ROWS[t]:
        put(ws, f"B{r_}", f'=IF({in_ref(cr, "A")}="","",{in_ref(cr, "A")})', f=font(10, False, INK), bg=WHITE,
            a=al("right", indent=1), b=Border(bottom=side(), right=side(accent, "thick")))
        for mi in range(12):
            col = L(MC1 + mi)
            put(ws, f"{L(3 + mi)}{r_}", f'=IF(OR({in_ref(cr, "A")}="",{in_ref(cr, col)}=""),"",{in_ref(cr, col)})',
                f=font(9, False, INK), bg=WHITE, fmt='#,##0;-#,##0;"–"', a=al(), b=Border(bottom=side(), left=side(WHITE)))
        put(ws, f"O{r_}", f'=IF(B{r_}="","",{in_ref(cr, "O")})', f=font(10, True, INK), bg="F8FAFC", fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
        ws.row_dimensions[r_].height = 20
        r_ += 1
    ws.conditional_formatting.add(
        f"C{first}:N{r_ - 1}",
        ColorScaleRule(start_type="min", start_color=WHITE, end_type="max", end_color=accent))
ws.freeze_panes = "A4"

# ================================================================== 🎯 יעדים
ws = ws_goal
setup(ws, "059669", {"A": 2, "B": 24, **{L(c): 13 for c in range(3, 12)}, "L": 2})
area(ws, 1, 60, 1, 12, CANVAS)
banner(ws, 2, 2, 11, "🎯 יעדי חיסכון וסילוק חובות",
       "ממלאים פעם אחת את התאים בצבע שמנת – ההתקדמות מחושבת לבד מתוך ההזנה החודשית")
ws.row_dimensions[1].height = 10

section_title(ws, "B5", "🐷 תוכניות חיסכון", TYPES["sav"][1])
header(ws, 6, 2, ["תוכנית", "יעד", "היה בצד בתחילת השנה", "תאריך יעד", "הופקד השנה", 'סה"כ נחסך',
                  "נותר", "התקדמות", "", "כמה להפקיד בחודש"], bg=TYPES["sav"][1])
ws.merge_cells("I6:J6")
S0 = 7
for i, cr in enumerate(CAT_ROWS["sav"]):
    r_ = S0 + i
    goal, start, date = SAVINGS_GOALS[i]
    put(ws, f"B{r_}", f'=IF({in_ref(cr, "A")}="","",{in_ref(cr, "A")})', f=font(10, True, INK), bg=WHITE,
        a=al("right", indent=1), b=Border(bottom=side(), right=side(TYPES["sav"][1], "thick")))
    put(ws, f"C{r_}", goal, f=font(10, False, INK), bg=INPUT_BG, fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
    put(ws, f"D{r_}", start, f=font(10, False, INK), bg=INPUT_BG, fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
    put(ws, f"E{r_}", "=DATE(" + ",".join(str(int(x)) for x in date.split("-")) + ")", f=font(10, False, INK), bg=INPUT_BG, fmt="mm/yyyy", a=al(), b=Border(bottom=side()))
    put(ws, f"F{r_}", f"={in_ref(cr, 'O')}", f=font(10, False, INK_2), bg=WHITE, fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
    put(ws, f"G{r_}", f"=N(D{r_})+F{r_}", f=font(10, True, INK), bg=WHITE, fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
    put(ws, f"H{r_}", f"=MAX(0,N(C{r_})-G{r_})", f=font(10, False, INK_2), bg=WHITE, fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
    put(ws, f"I{r_}", f'=IF(N(C{r_})=0,"",MIN(1,G{r_}/C{r_}))', f=font(11, True, TYPES["sav"][1]), bg=WHITE, fmt="0%", a=al(), b=Border(bottom=side()))
    put(ws, f"J{r_}", f'=IF(I{r_}="","",{bar_formula(f"I{r_}", 10)})', f=font(9, False, TYPES["sav"][1]), bg=WHITE, a=al(), b=Border(bottom=side()))
    months_left = f"MAX(1,(YEAR(E{r_})-YEAR(TODAY()))*12+MONTH(E{r_})-MONTH(TODAY()))"
    put(ws, f"K{r_}", f'=IF(OR(N(C{r_})=0,E{r_}=""),"",IF(H{r_}=0,"✅ הושג",H{r_}/{months_left}))',
        f=font(10, True, INK), bg=WHITE, fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
    ws.row_dimensions[r_].height = 24
S1 = S0 + len(CAT_ROWS["sav"]) - 1
ST = S1 + 1
put(ws, f"B{ST}", 'סה"כ', f=font(10, True, WHITE), bg=NAVY, a=al("right", indent=1))
for col in "CDFGH":
    put(ws, f"{col}{ST}", f"=SUM({col}{S0}:{col}{S1})", f=font(10, True, WHITE), bg=NAVY, fmt=CUR_PLAIN, a=al())
put(ws, f"I{ST}", f'=IF(C{ST}=0,"",MIN(1,G{ST}/C{ST}))', f=font(11, True, WHITE), bg=NAVY, fmt="0%", a=al())
put(ws, f"J{ST}", f'=IF(I{ST}="","",{bar_formula(f"I{ST}", 10)})', f=font(9, False, "A7F3D0"), bg=NAVY, a=al())
put(ws, f"K{ST}", f"=SUM(K{S0}:K{S1})", f=font(10, True, WHITE), bg=NAVY, fmt=CUR_PLAIN, a=al())
put(ws, f"E{ST}", None, bg=NAVY)
ws.row_dimensions[ST].height = 26

D_T = ST + 3
section_title(ws, f"B{D_T}", "💳 סילוק חובות", TYPES["debt"][1])
header(ws, D_T + 1, 2, ["חוב", "יתרה בתחילת השנה", "ריבית שנתית", "שולם השנה", "יתרה משוערת",
                        "נפרע", "", "החזר חודשי ממוצע", "חודשים לסיום", "סיום משוער"], bg=TYPES["debt"][1])
ws.merge_cells(f"G{D_T + 1}:H{D_T + 1}")
D0 = D_T + 2
for i, cr in enumerate(CAT_ROWS["debt"]):
    r_ = D0 + i
    bal, rate = DEBT_SETUP[i]
    put(ws, f"B{r_}", f'=IF({in_ref(cr, "A")}="","",{in_ref(cr, "A")})', f=font(10, True, INK), bg=WHITE,
        a=al("right", indent=1), b=Border(bottom=side(), right=side(TYPES["debt"][1], "thick")))
    put(ws, f"C{r_}", bal, f=font(10, False, INK), bg=INPUT_BG, fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
    put(ws, f"D{r_}", rate, f=font(10, False, INK), bg=INPUT_BG, fmt="0.0%", a=al(), b=Border(bottom=side()))
    put(ws, f"E{r_}", f"={in_ref(cr, 'O')}", f=font(10, False, INK_2), bg=WHITE, fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
    put(ws, f"F{r_}", f"=MAX(0,N(C{r_})-E{r_})", f=font(10, True, INK), bg=WHITE, fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
    put(ws, f"G{r_}", f'=IF(N(C{r_})=0,"",MIN(1,E{r_}/C{r_}))', f=font(11, True, TYPES["debt"][1]), bg=WHITE, fmt="0%", a=al(), b=Border(bottom=side()))
    put(ws, f"H{r_}", f'=IF(G{r_}="","",{bar_formula(f"G{r_}", 10)})', f=font(9, False, TYPES["debt"][1]), bg=WHITE, a=al(), b=Border(bottom=side()))
    put(ws, f"I{r_}", f"={in_ref(cr, 'P')}", f=font(10, False, INK_2), bg=WHITE, fmt=CUR_PLAIN, a=al(), b=Border(bottom=side()))
    put(ws, f"J{r_}", f'=IF(OR(N(C{r_})=0,N(I{r_})=0),"",IF(F{r_}=0,0,IFERROR(ROUNDUP(NPER(N(D{r_})/12,-I{r_},F{r_}),0),"ההחזר נמוך מהריבית")))',
        f=font(10, False, INK), bg=WHITE, fmt="0", a=al(), b=Border(bottom=side()))
    put(ws, f"K{r_}", f'=IF(ISNUMBER(J{r_}),IF(J{r_}=0,"🎉 נסגר",EDATE(TODAY(),J{r_})),"")',
        f=font(10, True, INK), bg=WHITE, fmt="mm/yyyy", a=al(), b=Border(bottom=side()))
    ws.row_dimensions[r_].height = 24
D1 = D0 + len(CAT_ROWS["debt"]) - 1
DT = D1 + 1
put(ws, f"B{DT}", 'סה"כ', f=font(10, True, WHITE), bg=NAVY, a=al("right", indent=1))
for col in "CEFI":
    put(ws, f"{col}{DT}", f"=SUM({col}{D0}:{col}{D1})", f=font(10, True, WHITE), bg=NAVY, fmt=CUR_PLAIN, a=al())
put(ws, f"G{DT}", f'=IF(C{DT}=0,"",MIN(1,E{DT}/C{DT}))', f=font(11, True, WHITE), bg=NAVY, fmt="0%", a=al())
put(ws, f"H{DT}", f'=IF(G{DT}="","",{bar_formula(f"G{DT}", 10)})', f=font(9, False, "FECACA"), bg=NAVY, a=al())
put(ws, f"J{DT}", f'=IF(COUNT(J{D0}:J{D1})=0,"",MAX(J{D0}:J{D1}))', f=font(10, True, WHITE), bg=NAVY, fmt="0", a=al())
put(ws, f"K{DT}", f'=IF(ISNUMBER(J{DT}),IF(J{DT}=0,"🎉",EDATE(TODAY(),J{DT})),"")', f=font(10, True, WHITE), bg=NAVY, fmt="mm/yyyy", a=al())
put(ws, f"D{DT}", None, bg=NAVY)
ws.row_dimensions[DT].height = 26
merge(ws, f"B{DT + 1}:K{DT + 1}",
      "💡 'יתרה משוערת' = יתרת תחילת השנה פחות ההחזרים (בלי ריבית). 'חודשים לסיום' כן מחשב ריבית לפי ההחזר החודשי הממוצע.",
      f=font(9, False, MUTED, italic=True), a=al("right", wrap=True))
ws.row_dimensions[DT + 1].height = 30

g1 = BarChart()
g1.type = "bar"
g1.add_data(Reference(ws, min_col=3, min_row=6, max_row=S1), titles_from_data=True)
g1.add_data(Reference(ws, min_col=7, min_row=6, max_row=S1), titles_from_data=True)
g1.set_categories(Reference(ws, min_col=2, min_row=S0, max_row=S1))
g1.series[0].graphicalProperties.solidFill = "D1FAE5"
g1.series[1].graphicalProperties.solidFill = TYPES["sav"][1]
for s in g1.series:
    s.graphicalProperties.line.noFill = True
g1.gapWidth = 40
g1.y_axis.numFmt = "#,##0"
g1.y_axis.majorGridlines = None
style_chart(g1, "יעד מול נחסך", 15.5, 8)
place(ws, g1, f"B{DT + 3}:E{DT + 19}")

g2 = BarChart()
g2.type = "bar"
g2.grouping = "stacked"
g2.overlap = 100
g2.add_data(Reference(ws, min_col=5, max_col=6, min_row=D_T + 1, max_row=D1), titles_from_data=True)
g2.set_categories(Reference(ws, min_col=2, min_row=D0, max_row=D1))
g2.series[0].graphicalProperties.solidFill = GOOD
g2.series[1].graphicalProperties.solidFill = "FCA5A5"
for s in g2.series:
    s.graphicalProperties.line.noFill = True
g2.gapWidth = 40
g2.y_axis.numFmt = "#,##0"
g2.y_axis.majorGridlines = None
style_chart(g2, "שולם מול נותר", 15.5, 8)
place(ws, g2, f"G{DT + 3}:K{DT + 19}")

# ================================================================== 📖 מדריך
ws = ws_guide
setup(ws, "94A3B8", {"A": 2, "B": 8, "C": 90, "D": 2}, landscape=False)
area(ws, 1, 40, 1, 4, CANVAS)
banner(ws, 2, 2, 3, "📖 איך משתמשים – 3 צעדים", "פעם אחת מגדירים, פעם בחודש ממלאים, וכל השאר קורה לבד")
ws.row_dimensions[1].height = 10
steps = [
    ("1", "הגדרה חד-פעמית (10 דקות)",
     f"בלשונית '{INP}': עדכנו את שמות הקטגוריות ואת התקציב החודשי לכל אחת (עמודה B). "
     f"בלשונית '{GOAL}': רשמו יעד לכל תוכנית חיסכון ואת יתרת החובות בתחילת השנה."),
    ("2", "פעם בחודש (5 דקות)",
     f"פותחים את דף הבנק ואת פירוט כרטיס האשראי, ובלשונית '{INP}' ממלאים את העמודה של החודש: "
     "כמה נכנס, כמה הופקד לחיסכון, כמה יצא על כל קטגוריה. מספרים בלבד, בלי ₪."),
    ("3", "מסתכלים על התמונה",
     f"'{DASH}' מציג אוטומטית את החודש האחרון שמילאתם: כרטיסי סיכום, תקציב מול בפועל, 5 ההוצאות הגדולות ותובנות. "
     f"אפשר לבחור כל חודש אחר מהרשימה. ב-'{TRD}' רואים את כל השנה, וב-'{GOAL}' את ההתקדמות ליעדים."),
]
r_ = 5
for num, head, body in steps:
    for rr in (r_, r_ + 1):
        for cc in (2, 3):
            ws.cell(rr, cc).fill = fill(WHITE)
    ws.merge_cells(f"B{r_}:B{r_ + 1}")
    put(ws, f"B{r_}", num, f=font(26, True, "F59E0B"), a=al())
    put(ws, f"C{r_}", head, f=font(13, True, INK), a=al("right", "bottom", indent=1))
    put(ws, f"C{r_ + 1}", body, f=font(10, False, INK_2), a=al("right", "top", wrap=True, indent=1))
    ws.cell(r_ + 1, 2).border = ws.cell(r_ + 1, 3).border = Border(bottom=side(CANVAS, "thick"))
    ws.row_dimensions[r_].height = 28
    ws.row_dimensions[r_ + 1].height = 48
    r_ += 2
r_ += 1
put(ws, f"C{r_}", "🎨 מקרא", f=font(12, True, INK), a=al("right"))
r_ += 1
legend = [(INPUT_BG, "תא להזנה – רק כאן מקלידים. כל השאר מחושב אוטומטית.")]
legend += [(TYPES[t][2], TYPES[t][0]) for t in ORDER]
legend += [("FEF2F2", "רקע אדום בלשונית ההזנה = חריגה מהתקציב של הקטגוריה")]
for color, text in legend:
    put(ws, f"B{r_}", None, bg=color, b=Border(*(side(LINE),) * 4))
    put(ws, f"C{r_}", text, f=font(10, False, INK), bg=WHITE, a=al("right", indent=1))
    ws.row_dimensions[r_].height = 22
    r_ += 1
r_ += 1
tips = [
    "💡 אין צורך למלא את כל הקטגוריות – קטגוריה ריקה פשוט לא נספרת.",
    "💡 רוצים קטגוריה חדשה? רשמו שם בשורה ריקה בסעיף המתאים – היא תופיע בכל הלשוניות.",
    "💡 לשנה חדשה: שכפלו את הקובץ, שנו את השנה בתא B3 בלשונית ההזנה ומחקו את המספרים של החודשים.",
]
for tip in tips:
    put(ws, f"C{r_}", tip, f=font(10, False, INK_2), a=al("right", wrap=True))
    ws.row_dimensions[r_].height = 22
    r_ += 1

wb.active = 0
wb.save(OUT)
print(f"saved {OUT}")
